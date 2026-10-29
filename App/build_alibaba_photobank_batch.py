import openpyxl
import sqlite3
import csv
import os
import glob

BASE_DIR = r"d:\Work\Hots\Product Input"
DB_PATH = os.path.join(BASE_DIR, "Database", "alibaba_master.db")
TEMPLATE_PATH = os.path.join(BASE_DIR, "Template", "Alibaba default template (recommended) (1).xlsx")
BATCH_DIR = os.path.join(BASE_DIR, "datagambar photbank")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "Output", "Alibaba_Bulk_Upload_November_2026.xlsx")
OUTPUT_CSV = os.path.join(BASE_DIR, "Output", "Alibaba_Bulk_Upload_November_2026.csv")

def load_all_photobank_uploaded_images():
    csv_files = sorted(glob.glob(os.path.join(BATCH_DIR, "*.csv")))
    all_uploaded = {}
    for cp in csv_files:
        with open(cp, 'r', encoding='utf-8') as f:
            for r in csv.DictReader(f):
                orig = r.get('Nama File Asli', '').strip()
                url = r.get('URL Gambar', '').strip()
                sku = r.get('SKU', '').strip()
                
                clean_name = orig
                if clean_name.startswith('APB-') and '_' in clean_name:
                    clean_name = clean_name.split('_', 1)[1]
                clean_name = clean_name.lstrip('_')
                
                all_uploaded[clean_name] = {
                    'clean_name': clean_name,
                    'url': url,
                    'photobank_sku': sku
                }
    return all_uploaded

def generate_b2b_export_keywords(title, cat_grp):
    cat_l = (cat_grp or "").lower()
    t_l = title.lower()

    if "safety" in cat_l or "s5" in t_l or "metaguard" in t_l or "max" in t_l:
        k1, k2, k3 = "steel toe safety boots", "s5 waterproof work boots", "chemical resistant industrial footwear"
    elif "orca" in cat_l or "orca" in t_l:
        k1, k2, k3 = "industrial rubber boots", "waterproof high safety boots", "heavy duty protective gumboots"
    elif "chef" in cat_l or "chef" in t_l:
        k1, k2, k3 = "slip resistant chef shoes", "kitchen work footwear", "waterproof food service clogs"
    elif "medic" in cat_l or "medic" in t_l:
        k1, k2, k3 = "medical safety clogs", "hospital cleanroom shoes", "anti slip nursing footwear"
    elif "allbike" in cat_l or "biker" in cat_l or "moto" in t_l or "riding" in t_l:
        k1, k2, k3 = "waterproof motorcycle riding shoes", "breathable pvc rain shoes", "casual outdoor commuter footwear"
    elif "terra" in cat_l or "terra" in t_l:
        k1, k2, k3 = "waterproof pvc rain boots", "heavy duty work gumboots", "agriculture plantation boots"
    elif "kids" in cat_l or "safari" in t_l or "2012" in t_l or "anak" in t_l:
        k1, k2, k3 = "kids waterproof rain boots", "anti slip children rubber boots", "durable cute toddler gumboots"
    elif "sandals" in cat_l or "clogs" in cat_l or "sandal" in t_l or "selop" in t_l:
        k1, k2, k3 = "comfortable slip on clogs", "waterproof garden slippers", "anti slip eva pvc slides"
    elif "sneakers" in cat_l or "star" in t_l:
        k1, k2, k3 = "casual waterproof sneakers", "pvc canvas style shoes", "daily lightweight rain footwear"
    elif any(k in t_l for k in ["bola", "963", "967", "969", "pantofel", "977"]):
        k1, k2, k3 = "waterproof molded cleat shoes", "durable pvc outdoor shoes", "slip resistant utility footwear"
    else:
        k1, k2, k3 = "waterproof pvc rain boots", "industrial protective footwear", "anti slip rubber work boots"

    return f"{k1}, {k2}, {k3}"

def main():
    print("=" * 70)
    print("      BUILDING ALIBABA NOVEMBER 2026 BULK UPLOAD BATCH")
    print("          (STRICT MATCHING: 21 PRODUK PHOTOBANK TERVERIFIKASI)")
    print("=" * 70)

    all_uploaded = load_all_photobank_uploaded_images()
    print(f"[*] Loaded {len(all_uploaded)} unique uploaded Photobank images from 'datagambar photbank/'.")

    # Strict definitions for 21 products
    exact_products_map = [
        {
            'db_sku': 'APB-001',
            'img_match': lambda name: any(k in name for k in ['allbike_orange_boot', 'orange_riding_shoes']) and 'allbike_2' not in name
        },
        {
            'db_sku': 'APB-013',
            'img_match': lambda name: 'allbike_2_orange' in name
        },
        {
            'db_sku': 'APB-097',
            'img_match': lambda name: name in ['ap_moto.webp', 'ap_moto_specs.webp', 'ap_boots_motorcycle_riding_shoes.webp']
        },
        {
            'db_sku': 'APB-098',
            'img_match': lambda name: 'ap_moto_2' in name
        },
        {
            'db_sku': 'APB-099',
            'img_match': lambda name: 'ap_moto_3' in name
        },
        {
            'db_sku': 'APB-015',
            'img_match': lambda name: 'terra_gp' in name
        },
        {
            'db_sku': 'APB-023',
            'img_match': lambda name: 'terra_1_putih' in name
        },
        {
            'db_sku': 'APB-022',
            'img_match': lambda name: ('terra_1' in name or 'ap_boots_terra_1' in name) and 'putih' not in name and 'gp' not in name
        },
        {
            'db_sku': 'APB-007',
            'img_match': lambda name: 'terra_eco_3' in name
        },
        {
            'db_sku': 'APB-025',
            'img_match': lambda name: 'metaguard' in name
        },
        {
            'db_sku': 'APB-024',
            'img_match': lambda name: 'orca' in name
        },
        {
            'db_sku': 'APB-030',
            'img_match': lambda name: 'chef' in name
        },
        {
            'db_sku': 'APB-035',
            'img_match': lambda name: ('blue_medical_boots' in name or 'ap_medic_-_banner' in name or 'ap medic (' in name)
        },
        {
            'db_sku': 'APB-037',
            'img_match': lambda name: 'ap_medic_-_white' in name
        },
        {
            'db_sku': 'APB-027',
            'img_match': lambda name: 'hobby' in name
        },
        {
            'db_sku': 'APB-006',
            'img_match': lambda name: ('2003' in name or 'triwarna' in name)
        },
        {
            'db_sku': 'APB-005',
            'img_match': lambda name: '2005' in name
        },
        {
            'db_sku': 'APB-002',
            'img_match': lambda name: ('2017' in name or 'green_boot' in name)
        },
        {
            'db_sku': 'APB-034',
            'img_match': lambda name: '9303' in name
        },
        {
            'db_sku': 'APB-028',
            'img_match': lambda name: '9506' in name
        },
        {
            'db_sku': 'APB-042',
            'img_match': lambda name: ('kids_safari' in name or 'safari_kids' in name or 'ap kids safari' in name)
        }
    ]

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    valid_batch_products = []

    for pdef in exact_products_map:
        sku = pdef['db_sku']
        matcher = pdef['img_match']
        
        cur.execute("SELECT * FROM products WHERE sku = ?", (sku,))
        p = cur.fetchone()
        if not p:
            print(f"[!] Warning: SKU {sku} not found in database!")
            continue
            
        prod = dict(p)
        title = prod['product_title']
        cat = prod['category_group']
        
        # Match only exact images
        matched_urls = []
        for clean_name, d in all_uploaded.items():
            if matcher(clean_name.lower()):
                u = d['url']
                if u not in matched_urls:
                    matched_urls.append(u)
                    
        if not matched_urls:
            print(f"[-] No images matched for {sku}: {title}")
            continue
            
        # Up to 6 unique images
        final_imgs = (matched_urls + [None] * 6)[:6]
        new_keywords = generate_b2b_export_keywords(title, cat)

        # Update in Database
        cur.execute("""
            UPDATE products
            SET image_1 = ?, image_2 = ?, image_3 = ?, image_4 = ?, image_5 = ?, image_6 = ?,
                keywords = ?, updated_at = CURRENT_TIMESTAMP
            WHERE sku = ?
        """, (
            final_imgs[0], final_imgs[1], final_imgs[2],
            final_imgs[3], final_imgs[4], final_imgs[5],
            new_keywords, sku
        ))

        p_up = prod.copy()
        for i in range(1, 7):
            p_up[f'image_{i}'] = final_imgs[i-1]
        p_up['keywords'] = new_keywords
        valid_batch_products.append(p_up)

    conn.commit()
    conn.close()

    print(f"\n[+] Total STRICT products selected for Alibaba Bulk Batch: {len(valid_batch_products)} products")

    # Populate into Excel Template
    print(f"[*] Loading Alibaba Excel template: {TEMPLATE_PATH}")
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb['Alibaba default template']

    # Delete existing data rows starting from row 5
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

    start_row = 5
    csv_rows = []

    header_row3 = [ws.cell(3, c).value for c in range(1, 58)]
    header_row4 = [ws.cell(4, c).value for c in range(1, 58)]
    csv_rows.append(header_row3)
    csv_rows.append(header_row4)

    for idx, prod in enumerate(valid_batch_products):
        row_idx = start_row + idx

        sku = prod['sku']
        title = prod['product_title']
        cat = prod['category']
        html_desc = prod['description_html']
        weight_kg = prod['gross_weight_kg'] or 1.7
        l_val = prod['pkg_length'] or 60.0
        w_val = prod['pkg_width'] or 42.0
        h_val = prod['pkg_height'] or 38.0

        t1_price = prod['tier1_price'] or prod['price_usd'] or 5.50
        t2_price = prod['tier2_price'] or round(t1_price * 0.95, 2)
        t3_price = prod['tier3_price'] or round(t1_price * 0.90, 2)
        t4_price = prod['tier4_price'] or round(t1_price * 0.85, 2)
        moq_p = prod['moq1'] or 100

        row_data = [
            title,                      # Col 1 (A): * Product title
            prod.get('image_1'),        # Col 2 (B): * Product image 1
            prod.get('image_2'),        # Col 3 (C): Product image 2
            prod.get('image_3'),        # Col 4 (D): Product image 3
            prod.get('image_4'),        # Col 5 (E): Product image 4
            prod.get('image_5'),        # Col 6 (F): Product image 5
            prod.get('image_6'),        # Col 7 (G): Product image 6
            html_desc,                  # Col 8 (H): Product description
            prod.get('place_of_origin') or "Indonesia", # Col 9 (I): Place of origin
            prod.get('brand_name') or "AP Boots",      # Col 10 (J): Brand name
            cat,                        # Col 11 (K): Category
            "Outsole Material",         # Col 12 (L): Product attribute name 1
            prod.get('outsole_material') or "PVC / Rubber", # Col 13 (M): Product attribute value 1
            "Upper Material",           # Col 14 (N): Product attribute name 2
            prod.get('upper_material') or "PVC",            # Col 15 (O): Product attribute value 2
            "Season",                   # Col 16 (P): Product attribute name 3
            prod.get('season') or "All Seasons",            # Col 17 (Q): Product attribute value 3
            "Feature",                  # Col 18 (R): Product attribute name 4
            prod.get('feature') or "Waterproof, Anti-Slip, High Durability", # Col 19 (S): Product attribute value 4
            "Gender",                   # Col 20 (T): Product attribute name 5
            prod.get('gender') or "Men / Unisex",           # Col 21 (U): Product attribute value 5
            prod.get('sell_by') or "piece",                 # Col 22 (V): Sell product by
            None,                       # Col 23 (W): Batch quantity
            prod.get('unit') or "Pair/Pairs",               # Col 24 (X): Unit
            prod.get('currency') or "USD",                  # Col 25 (Y): Currency
            prod.get('pricing_type') or "Tiered pricing",   # Col 26 (Z): Pricing type
            prod.get('inventory') or 5000,                  # Col 27 (AA): Inventory
            sku,                        # Col 28 (AB): SPU ID (Model number)
            sku,                        # Col 29 (AC): SKU code
            t1_price,                   # Col 30 (AD): SKU pricing
            "Color",                    # Col 31 (AE): SKU attribute name 1
            prod.get('color') or "Black; Green; White; Yellow", # Col 32 (AF): SKU attribute value 1
            "Size",                     # Col 33 (AG): SKU attribute name 2
            prod.get('size_range') or "EU 38-44",              # Col 34 (AH): SKU attribute value 2
            None,                       # Col 35 (AI): SKU attribute name 3
            None,                       # Col 36 (AJ): SKU attribute value 3
            None,                       # Col 37 (AK): SKU attribute name 4
            None,                       # Col 38 (AL): SKU attribute value 4
            t1_price,                   # Col 39 (AM): Tiered pricing 1
            moq_p,                      # Col 40 (AN): MOQ 1
            t2_price,                   # Col 41 (AO): Tiered pricing 2
            moq_p * 5,                  # Col 42 (AP): MOQ 2
            t3_price,                   # Col 43 (AQ): Tiered pricing 3
            moq_p * 10,                 # Col 44 (AR): MOQ 3
            t4_price,                   # Col 45 (AS): Range pricing (Lower limit)
            t1_price,                   # Col 46 (AT): Range pricing (Higher limit)
            weight_kg,                  # Col 47 (AU): * Gross weight (KG)
            prod.get('shipping_template') or "Default Shipping Template", # Col 48 (AV): * Shipping template name
            moq_p,                      # Col 49 (AW): * Shipping quantity 1
            prod.get('lead_time_1') or 7, # Col 50 (AX): * Estimated shipping lead time 1
            moq_p * 5,                  # Col 51 (AY): Shipping quantity 2
            prod.get('lead_time_2') or 14, # Col 52 (AZ): Estimated shipping lead time 2
            moq_p * 10,                 # Col 53 (BA): Shipping quantity 3
            prod.get('lead_time_3') or 30, # Col 54 (BB): Estimated shipping lead time 3
            l_val,                      # Col 55 (BC): Product size-Length
            w_val,                      # Col 56 (BD): Product size-Width
            h_val                       # Col 57 (BE): Product size-Height
        ]

        for col_i, val in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_i, value=val)

        csv_rows.append(row_data)

    print(f"[*] Saving November Batch Excel to: {OUTPUT_EXCEL}")
    wb.save(OUTPUT_EXCEL)

    with open(OUTPUT_CSV, 'w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerows(csv_rows)
    print(f"[*] Saving November Batch CSV to: {OUTPUT_CSV}")

    print("=" * 70)
    print(f"SUCCESS: {len(valid_batch_products)} products generated with 100% STRICT matching Photobank URLs!")
    print("=" * 70)

if __name__ == '__main__':
    main()
