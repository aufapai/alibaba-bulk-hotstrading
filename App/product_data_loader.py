import os
import re
import csv
import io
import sqlite3
import openpyxl
import pandas as pd
from datetime import datetime

try:
    from image_deduplicator import clean_and_deduplicate_images
    from keyword_bot import match_product_keywords
except ImportError:
    from App.image_deduplicator import clean_and_deduplicate_images
    from App.keyword_bot import match_product_keywords

BASE_DIR = r"d:\Work\Hots\Product Input"
DB_PATH = os.path.join(BASE_DIR, "Database", "alibaba_master.db") if os.path.exists(os.path.join(BASE_DIR, "Database", "alibaba_master.db")) else os.path.join(BASE_DIR, "alibaba_master.db")
TEMPLATE_PATH = os.path.join(BASE_DIR, "Template", "Alibaba default template (recommended) (1).xlsx") if os.path.exists(os.path.join(BASE_DIR, "Template", "Alibaba default template (recommended) (1).xlsx")) else os.path.join(BASE_DIR, "Alibaba default template (recommended) (1).xlsx")
EXPORT_PATH_XLSX = os.path.join(BASE_DIR, "Output", "Alibaba_Bulk_Upload_Tokopedia_Master_2026.xlsx")
EXPORT_PATH_CSV = os.path.join(BASE_DIR, "Output", "Alibaba_Bulk_Upload_Tokopedia_Master_2026.csv")

# 57 Standard Alibaba Column Headers (Row 3 in template)
ALIBABA_57_COLUMNS = [
    '* Product title', '* Product image 1', 'Product image 2', 'Product image 3',
    'Product image 4', 'Product image 5', 'Product image 6', 'Product description',
    'Place of origin', 'Brand name', 'Category', 'Product attribute name 1',
    'Product attribute value 1', 'Product attribute name 2', 'Product attribute value 2',
    'Product attribute name 3', 'Product attribute value 3', 'Product attribute name 4',
    'Product attribute value 4', 'Product attribute name 5', 'Product attribute value 5',
    'Sell product by', 'Batch quantity', 'Unit', 'Currency', 'Pricing type',
    'Inventory', 'SPU ID (Model number)', 'SKU code', 'SKU pricing',
    'SKU attribute name 1', 'SKU attribute value 1', 'SKU attribute name 2',
    'SKU attribute value 2', 'SKU attribute name 3', 'SKU attribute value 3',
    'SKU attribute name 4', 'SKU attribute value 4', 'Tiered pricing 1', 'MOQ 1',
    'Tiered pricing 2', 'MOQ 2', 'Tiered pricing 3', 'MOQ 3',
    'Range pricing (Lower limit)', 'Range pricing (Higher limit)', '* Gross weight (KG)',
    '* Shipping template name', '* Shipping quantity 1', '* Estimated shipping lead time 1',
    'Shipping quantity 2', 'Estimated shipping lead time 2', 'Shipping quantity 3',
    'Estimated shipping lead time 3', 'Product size-Length', 'Product size-Width',
    'Product size-Height'
]

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def get_existing_image_map():
    """Build a mapping of SKU / Model / Product Name keywords to 6 CDN image URLs."""
    img_map = {}
    if not os.path.exists(DB_PATH):
        return img_map
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT sku, model_number, product_title, image_1, image_2, image_3, image_4, image_5, image_6 FROM products WHERE image_1 IS NOT NULL AND image_1 != ''")
        for r in cur.fetchall():
            imgs = [r['image_1'], r['image_2'] or '', r['image_3'] or '', r['image_4'] or '', r['image_5'] or '', r['image_6'] or '']
            if r['sku']:
                img_map[str(r['sku']).strip().upper()] = imgs
                # Also map base prefix if any (e.g., APB-001 or AP-MX-01)
                prefix = str(r['sku']).strip().upper().split('-V')[0]
                img_map[prefix] = imgs
            if r['model_number']:
                img_map[str(r['model_number']).strip().upper()] = imgs
            if r['product_title']:
                clean_title = re.sub(r'[^A-Z0-9]', '', str(r['product_title']).upper())[:25]
                img_map[clean_title] = imgs
        conn.close()
    except Exception as e:
        print(f"[WARN] Error loading image map: {e}")
    return img_map

def determine_category_group(title, category=""):
    t_low = (title or "").lower()
    c_low = (category or "").lower()
    combined = t_low + " " + c_low

    if "terra" in combined: return "AP Terra Series"
    elif "orca" in combined: return "AP Orca Series"
    elif "chef" in combined: return "AP Chef Series"
    elif "medic" in combined: return "AP Medic Series"
    elif "safety" in combined or "s5" in combined or "s4" in combined or "steel toe" in combined: return "AP Safety Boots"
    elif "star" in combined: return "AP Kasual Sneakers"
    elif "allbike" in combined or "moto" in combined or "biker" in combined or "hobby" in combined: return "AP Biker & Casual"
    elif "kids" in combined or "anak" in combined: return "AP Kids Boots"
    elif "clogs" in combined or "sandal" in combined or "995" in combined: return "AP Sandals & Clogs"
    else: return "AP Classic Boots"

def generate_keywords(title, category_grp):
    cat_l = (category_grp or "").lower()
    t_l = (title or "").lower()
    if "terra" in cat_l or "terra" in t_l:
        return "waterproof pvc rain boots, agriculture safety boots, heavy duty work shoes"
    elif "orca" in cat_l or "orca" in t_l:
        return "industrial rain boots, protective high boots, anti slip rubber boots"
    elif "chef" in cat_l or "chef" in t_l:
        return "slip resistant chef shoes, kitchen work footwear, waterproof slip on clogs"
    elif "medic" in cat_l or "medic" in t_l:
        return "medical safety shoes, cleanroom clogs, hospital anti slip footwear"
    elif "safety" in cat_l or "safety" in t_l:
        return "steel toe safety boots, s5 waterproof work boots, chemical resistant footwear"
    elif "allbike" in cat_l or "moto" in t_l:
        return "motorcycle riding shoes, waterproof rain shoes, breathable outdoor footwear"
    elif "star" in cat_l or "star" in t_l:
        return "casual waterproof sneakers, pvc canvas style shoes, daily rain kets"
    elif "kids" in cat_l or "kids" in t_l:
        return "kids rain boots, waterproof children boots, toddler cute rubber boots"
    elif "sandals" in cat_l or "clogs" in cat_l:
        return "anti slip work slippers, waterproof garden clogs, eva pvc slides"
    else:
        return "waterproof rain boots, safety work footwear, made in indonesia boots"

def list_available_files():
    """List all available XLSX and CSV files in Input/ and Output/ folders."""
    files = []
    folders = [
        ("Input", os.path.join(BASE_DIR, "Input")),
        ("Output", os.path.join(BASE_DIR, "Output"))
    ]
    for folder_name, folder_path in folders:
        if not os.path.exists(folder_path):
            continue
        for fname in os.listdir(folder_path):
            ext = os.path.splitext(fname)[1].lower()
            if ext in ['.xlsx', '.xls', '.csv']:
                fpath = os.path.join(folder_path, fname)
                try:
                    stat = os.stat(fpath)
                    files.append({
                        "id": f"{folder_name}/{fname}",
                        "name": fname,
                        "folder": folder_name,
                        "format": ext.replace('.', '').upper(),
                        "size_kb": round(stat.st_size / 1024, 1),
                        "modified": datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M'),
                        "path": fpath
                    })
                except Exception:
                    pass
    # Sort files by modified descending
    files.sort(key=lambda x: x['modified'], reverse=True)
    return files

def read_raw_data(filepath, file_content=None, filename=""):
    """Reads either a filepath or raw bytes/string into a DataFrame and schema hint."""
    ext = os.path.splitext(filename or filepath or "")[1].lower()
    if not ext and filepath:
        ext = os.path.splitext(filepath)[1].lower()

    if file_content is not None:
        if isinstance(file_content, str):
            if ';base64,' in file_content:
                file_content = file_content.split(';base64,')[1]
            try:
                import base64
                decoded = base64.b64decode(file_content)
                file_obj = io.BytesIO(decoded)
            except Exception:
                file_obj = io.StringIO(file_content)
        elif isinstance(file_content, bytes):
            file_obj = io.BytesIO(file_content)
        else:
            file_obj = file_content
    else:
        file_obj = filepath

    df = None
    sheet_name = None

    if ext in ['.xlsx', '.xls']:
        xl = pd.ExcelFile(file_obj)
        # Select best sheet
        if '800_Listings_Master' in xl.sheet_names:
            sheet_name = '800_Listings_Master'
            df = xl.parse(sheet_name)
        elif 'Alibaba default template' in xl.sheet_names:
            sheet_name = 'Alibaba default template'
            # Check row 2 vs row 0
            df_probe = xl.parse(sheet_name, nrows=5, header=None)
            # Find row containing '* Product title'
            title_row = 0
            for r_idx in range(len(df_probe)):
                row_vals = [str(x) for x in df_probe.iloc[r_idx].values if pd.notna(x)]
                if any('* Product title' in x or 'Product title' in x for x in row_vals):
                    title_row = r_idx
                    break
            df = xl.parse(sheet_name, header=title_row)
        else:
            sheet_name = xl.sheet_names[0]
            df = xl.parse(sheet_name)
    else:
        # CSV parsing
        try:
            # First try auto-sniff with python engine
            df = pd.read_csv(file_obj, sep=None, engine='python', encoding='utf-8')
        except Exception:
            if hasattr(file_obj, 'seek'): file_obj.seek(0)
            try:
                df = pd.read_csv(file_obj, sep=',', encoding='utf-8-sig')
            except Exception:
                if hasattr(file_obj, 'seek'): file_obj.seek(0)
                df = pd.read_csv(file_obj, sep=';', encoding='latin1')

    # Clean header: strip strings and filter out non-data guideline row if present
    df.columns = [str(c).strip() for c in df.columns]

    # If first row contains guidelines (e.g., '1. Do not enter any unsupported symbols'), drop it
    if len(df) > 0:
        first_row_vals = " ".join([str(x) for x in df.iloc[0].values[:5] if pd.notna(x)])
        if "1. Do not enter" in first_row_vals or "Same rule as cell" in first_row_vals or "One cell for one" in first_row_vals:
            df = df.iloc[1:].reset_index(drop=True)

    return df, ext, sheet_name

def parse_records(df, preserve_images=True):
    """Normalizes DataFrame rows into standardized product dictionaries for DB upsert."""
    records = []
    cols = list(df.columns)
    img_map = get_existing_image_map() if preserve_images else {}

    # Helper function to find column value by alternative names
    def get_val(row, *candidates, default=''):
        for c in candidates:
            for col in cols:
                if col.lower() == c.lower() or col.lower().startswith(c.lower()):
                    v = row.get(col)
                    if pd.notna(v) and str(v).strip() != '' and str(v).strip().lower() != 'nan':
                        return v
        return default

    # Detect Schema Type
    is_alibaba_template = any('* Product title' in c or 'SPU ID (Model number)' in c or 'Tiered pricing 1' in c for c in cols)
    is_800_matrix = any('Unique_SKU_Code' in c or 'Target_Industry_Angle' in c or 'Base_SKU' in c for c in cols)
    is_lazada = any('lazada' in c.lower() for c in cols) or ('rating_score' in cols and 'item_page_link' in cols) or ('quantity_input' in cols and 'shipping_fee' in cols)
    is_blibli = any('blibli' in c.lower() for c in cols) or ('guaranteed_original' in cols and 'free_return' in cols) or ('sales_count' in cols and any('image2' in c.lower() for c in cols))

    for idx, row in df.iterrows():
        sku = ""
        model_number = ""
        title = ""
        category = ""
        price_usd = 5.50
        tier1_price = 5.50
        tier2_price = None
        tier3_price = None
        tier4_price = None
        moq1 = 100
        moq2 = 500
        moq3 = 1000
        gross_weight = 1.7
        imgs = ['', '', '', '', '', '']
        desc_html = ""
        color = "Black; Green; White; Yellow"
        size_range = "EU 38-44"
        outsole = "PVC / Rubber"
        upper = "PVC"
        season = "All Seasons"
        feature = "100% Waterproof, Anti-Slip"
        gender = "Men / Unisex"
        brand = "AP Boots"
        origin = "Indonesia"
        keywords = ""
        lead1 = 7
        lead2 = 14
        lead3 = 30

        if is_alibaba_template:
            title = str(get_val(row, '* Product title', 'Product title', default='AP Boots')).strip()
            imgs = [
                str(get_val(row, '* Product image 1', 'Product image 1', default='')),
                str(get_val(row, 'Product image 2', default='')),
                str(get_val(row, 'Product image 3', default='')),
                str(get_val(row, 'Product image 4', default='')),
                str(get_val(row, 'Product image 5', default='')),
                str(get_val(row, 'Product image 6', default='')),
            ]
            desc_html = str(get_val(row, 'Product description', default=''))
            origin = str(get_val(row, 'Place of origin', default='Indonesia'))
            brand = str(get_val(row, 'Brand name', default='AP Boots'))
            category = str(get_val(row, 'Category', default="Shoes & Accessories > Men's Shoes > Boots"))
            outsole = str(get_val(row, 'Product attribute value 1', default='PVC / Rubber'))
            upper = str(get_val(row, 'Product attribute value 2', default='PVC'))
            season = str(get_val(row, 'Product attribute value 3', default='All Seasons'))
            feature = str(get_val(row, 'Product attribute value 4', default='100% Waterproof, Anti-Slip'))
            gender = str(get_val(row, 'Product attribute value 5', default='Men / Unisex'))
            model_number = str(get_val(row, 'SPU ID (Model number)', default=f'APB-{idx+1:03d}')).strip()
            sku = str(get_val(row, 'SKU code', default=model_number)).strip()
            try: price_usd = float(get_val(row, 'SKU pricing', default=5.50))
            except: price_usd = 5.50
            color = str(get_val(row, 'SKU attribute value 1', default='Black; Green; White; Yellow'))
            size_range = str(get_val(row, 'SKU attribute value 2', default='EU 38-44'))
            try: tier1_price = float(get_val(row, 'Tiered pricing 1', default=price_usd))
            except: tier1_price = price_usd
            try: moq1 = int(get_val(row, 'MOQ 1', default=100))
            except: moq1 = 100
            try: tier2_price = float(get_val(row, 'Tiered pricing 2', default=round(tier1_price*0.95, 2)))
            except: tier2_price = None
            try: moq2 = int(get_val(row, 'MOQ 2', default=500))
            except: moq2 = 500
            try: tier3_price = float(get_val(row, 'Tiered pricing 3', default=round(tier1_price*0.90, 2)))
            except: tier3_price = None
            try: moq3 = int(get_val(row, 'MOQ 3', default=1000))
            except: moq3 = 1000
            try: gross_weight = float(get_val(row, '* Gross weight (KG)', 'Gross weight (KG)', default=1.7))
            except: gross_weight = 1.7

        elif is_800_matrix:
            sku = str(get_val(row, 'Unique_SKU_Code', default=f'APB-V{idx+1}')).strip()
            model_number = str(get_val(row, 'Base_SKU', default=sku.split('-V')[0])).strip()
            title = str(get_val(row, 'Alibaba_Product_Title', 'Base_Product_Name', default='AP Boots')).strip()
            category = str(get_val(row, 'Alibaba_Category', default="Shoes & Accessories > Men's Shoes > Boots")).strip()
            brand = str(get_val(row, 'Brand', default='AP Boots')).strip()
            origin = str(get_val(row, 'Place_of_Origin', default='Surabaya, Indonesia')).strip()
            upper = str(get_val(row, 'Upper_Material', default='PVC / Synthetic Rubber Blend')).strip()
            outsole = str(get_val(row, 'Outsole_Material', default='Nitrile Rubber / Heavy Duty PVC')).strip()
            color = str(get_val(row, 'Colors', default='Black; Green; White; Yellow')).strip()
            size_range = str(get_val(row, 'Size_Range', default='EU 38 - 43')).strip()
            feature = str(get_val(row, 'Key_Features', default='Waterproof, Anti-Slip, Oil Resistant')).strip()
            try: moq1 = int(get_val(row, 'MOQ_Pairs', default=60))
            except: moq1 = 60
            try:
                tier1_price = float(get_val(row, 'Price_Tier1_USD (MOQ 60)', 'Price_Tier1_USD', default=5.50))
            except: tier1_price = 5.50
            price_usd = tier1_price
            try: tier2_price = float(get_val(row, 'Price_Tier2_USD (MOQ 300)', default=round(tier1_price*0.85, 2)))
            except: tier2_price = round(tier1_price*0.85, 2)
            try: tier3_price = float(get_val(row, 'Price_Tier3_USD (MOQ 1200)', default=round(tier1_price*0.70, 2)))
            except: tier3_price = round(tier1_price*0.70, 2)
            moq2 = 300
            moq3 = 1200
            try:
                gw = float(get_val(row, 'Carton_Gross_Weight_KG', default=18.0))
                gross_weight = round(gw / 12.0, 2) # Per pair weight
            except: gross_weight = 1.7
            desc_html = str(get_val(row, 'Tailored_HTML_Description', default=''))
            kw_parts = [
                str(get_val(row, 'Primary_Keyword_1', default='')),
                str(get_val(row, 'Primary_Keyword_2', default='')),
                str(get_val(row, 'Primary_Keyword_3', default='')),
                str(get_val(row, 'Longtail_Keywords', default=''))
            ]
            keywords = ", ".join([k.strip() for k in kw_parts if k.strip()])

        elif is_lazada:
            sku = f"LAZ-{idx+1:03d}"
            title = str(get_val(row, 'name', 'title', default='AP Boots Lazada')).strip()
            model_number = title.split('-')[0].strip()[:30]
            # Price parsing
            data_str = str(get_val(row, 'data', default=''))
            price_idr = None
            if 'Rp' in data_str:
                m_p = re.search(r'Rp\s*([\d\.]+)', data_str)
                if m_p:
                    price_idr = int(m_p.group(1).replace('.', ''))
            if not price_idr and pd.notna(row.get('price')):
                try:
                    p = float(row.get('price'))
                    price_idr = int(p * 1000) if p < 1000 else int(p)
                except:
                    price_idr = 50000
            price_usd = round((price_idr or 50000) / 16000.0, 2)
            tier1_price = price_usd
            tier2_price = round(price_usd * 0.95, 2)
            tier3_price = round(price_usd * 0.90, 2)
            gross_weight = 1.7
            # Deduplicate images
            raw_laz_imgs = str(get_val(row, 'image_1', default='')) + '\n' + str(get_val(row, 'image_27', default=''))
            imgs = clean_and_deduplicate_images(raw_laz_imgs.split(), marketplace='lazada')
            # Match 800 matrix keywords
            kw_info = match_product_keywords(title, category="Boots")
            keywords = kw_info['keywords']
            title = kw_info['alibaba_product_title']
            desc_html = f"<p>Official AP Boots - {title}. Available for B2B export supply.</p>"

        elif is_blibli:
            sku = f"BLI-{idx+1:03d}"
            title = str(get_val(row, 'name', 'item_page_title', default='AP Boots Blibli')).strip()
            model_number = title.split('-')[0].strip()[:30]
            # Price parsing (decimal thousands e.g. 120.0 -> 120000)
            raw_val = str(get_val(row, 'data', 'price', default=''))
            m_b = re.search(r'([\d\.]+)', raw_val)
            price_idr = None
            if m_b:
                val_f = float(m_b.group(1))
                price_idr = int(val_f * 1000) if val_f < 1000 else int(val_f)
            price_usd = round((price_idr or 75000) / 16000.0, 2)
            tier1_price = price_usd
            tier2_price = round(price_usd * 0.95, 2)
            tier3_price = round(price_usd * 0.90, 2)
            gross_weight = 1.7
            # Images
            raw_bli_imgs = []
            for c in cols:
                if 'image' in c.lower():
                    v = str(row.get(c) or '').strip()
                    if v.startswith('http'):
                        raw_bli_imgs.append(v)
            imgs = clean_and_deduplicate_images(raw_bli_imgs, marketplace='blibli')
            # Match 800 matrix keywords
            kw_info = match_product_keywords(title, category="Boots")
            keywords = kw_info['keywords']
            title = kw_info['alibaba_product_title']
            desc_html = f"<p>Official AP Boots - {title}. Premium quality Indonesian footwear.</p>"

        else:
            # Generic CSV / XLSX
            sku = str(get_val(row, 'sku', 'SKU', 'model_number', default=f'APB-{idx+1:03d}')).strip()
            model_number = str(get_val(row, 'model_number', 'model', 'Base_SKU', default=sku)).strip()
            title = str(get_val(row, 'product_name', 'product_title', 'name', 'title', default='AP Boots Product')).strip()
            category = str(get_val(row, 'category', 'category_src', 'category_source', default="Shoes & Accessories > Men's Shoes > Boots")).strip()
            brand = str(get_val(row, 'brand_name', 'brand', default='AP Boots')).strip()
            origin = str(get_val(row, 'place_of_origin', 'origin', default='Indonesia')).strip()
            outsole = str(get_val(row, 'outsole_material', 'outsole', default='PVC / Rubber')).strip()
            upper = str(get_val(row, 'upper_material', 'upper', default='PVC')).strip()
            color = str(get_val(row, 'color', 'colors', 'warna', default='Black; Green; White; Yellow')).strip()
            size_range = str(get_val(row, 'size', 'size_range', 'ukuran', default='EU 38-44')).strip()
            feature = str(get_val(row, 'feature', 'features', default='Waterproof, Anti-Slip')).strip()
            desc_html = str(get_val(row, 'description_html', 'description', 'deskripsi', default=''))

            # Price parsing
            raw_price = get_val(row, 'base_usd', 'price_usd', 'price', 'tier1_usd', default='')
            if not raw_price:
                raw_idr = get_val(row, 'harga_idr', 'price_idr', default='')
                try: price_usd = round(float(raw_idr) / 16000.0, 2)
                except: price_usd = 5.50
            else:
                try: price_usd = float(raw_price)
                except: price_usd = 5.50
            tier1_price = price_usd

            try: tier2_price = float(get_val(row, 'tier2_usd', 'tier2_price', default=round(price_usd*0.95, 2)))
            except: tier2_price = round(price_usd*0.95, 2)
            try: tier3_price = float(get_val(row, 'tier3_usd', 'tier3_price', default=round(price_usd*0.90, 2)))
            except: tier3_price = round(price_usd*0.90, 2)
            try: tier4_price = float(get_val(row, 'tier4_usd', 'tier4_price', default=round(price_usd*0.85, 2)))
            except: tier4_price = round(price_usd*0.85, 2)

            try: gross_weight = float(get_val(row, 'weight_kg', 'gross_weight_kg', default=1.7))
            except: gross_weight = 1.7

            # Image resolution
            img_c = str(get_val(row, 'product_photo', 'image_1', 'google_img_url', 'image_local', default=''))
            imgs[0] = img_c

        # Clean SKU and Title
        if not sku:
            sku = f"APB-{idx+1:03d}"
        if not title:
            title = f"AP Boots {sku}"

        # Clean and fallback images from cache / DB if empty
        if preserve_images and (not imgs[0] or not imgs[0].startswith('http')):
            # Try to match in image map
            matched_imgs = (
                img_map.get(sku.upper()) or
                img_map.get(model_number.upper()) or
                img_map.get(sku.split('-V')[0].upper()) or
                img_map.get(model_number.split('-V')[0].upper())
            )
            if not matched_imgs:
                # Match title prefix
                clean_t = re.sub(r'[^A-Z0-9]', '', title.upper())[:25]
                matched_imgs = img_map.get(clean_t)
            if matched_imgs:
                imgs = list(matched_imgs)

        # Fallback default image if still empty
        if not imgs[0] or not imgs[0].startswith('http'):
            imgs[0] = "https://p16-images-comn-sg.tokopedia-static.net/tos-alisg-i-aphluv4xwc-sg/c134f882bb9646459cfbb06e2a19ddb0~tplv-aphluv4xwc-white-pad-v1:1600:1600.jpeg"

        # Categorization & Keywords
        cat_group = determine_category_group(title, category)
        if not keywords:
            keywords = generate_keywords(title, cat_group)

        # HTML Description Fallback
        if not desc_html or len(desc_html.strip()) < 20:
            desc_html = f"""<div style="font-family: Arial, sans-serif; line-height: 1.6; color: #333333; max-width: 900px;">
  <div style="background: linear-gradient(135deg, #1e3a8a, #2563eb); color: #ffffff; padding: 24px; border-radius: 8px;">
    <h2 style="margin: 0; font-size: 24px;">{title}</h2>
    <p style="margin: 8px 0 0; opacity: 0.9;">Official B2B Export Supply — PT Indo Sukses Semesta (AP Boots Indonesia)</p>
  </div>
  <div style="padding: 20px 0;">
    <h3>Key Specifications</h3>
    <ul>
      <li><b>Brand:</b> {brand}</li>
      <li><b>Country of Origin:</b> {origin}</li>
      <li><b>Upper Material:</b> {upper}</li>
      <li><b>Outsole Material:</b> {outsole}</li>
      <li><b>Available Sizes:</b> {size_range}</li>
      <li><b>Color Variants:</b> {color}</li>
      <li><b>Special Features:</b> {feature}</li>
    </ul>
  </div>
</div>"""

        records.append({
            "sku": sku,
            "model_number": model_number or sku,
            "product_title": title,
            "category": category or "Shoes & Accessories > Men's Shoes > Boots",
            "category_group": cat_group,
            "keywords": keywords,
            "image_1": imgs[0],
            "image_2": imgs[1] if len(imgs) > 1 else '',
            "image_3": imgs[2] if len(imgs) > 2 else '',
            "image_4": imgs[3] if len(imgs) > 3 else '',
            "image_5": imgs[4] if len(imgs) > 4 else '',
            "image_6": imgs[5] if len(imgs) > 5 else '',
            "description_html": desc_html,
            "place_of_origin": origin,
            "brand_name": brand,
            "outsole_material": outsole,
            "upper_material": upper,
            "season": season,
            "feature": feature,
            "gender": gender,
            "sell_by": "piece",
            "unit": "Pair/Pairs",
            "currency": "USD",
            "pricing_type": "Tiered pricing",
            "inventory": 5000,
            "price_usd": round(price_usd, 2),
            "color": color,
            "size_range": size_range,
            "tier1_price": round(tier1_price, 2) if tier1_price else round(price_usd, 2),
            "moq1": moq1,
            "tier2_price": round(tier2_price, 2) if tier2_price else None,
            "moq2": moq2,
            "tier3_price": round(tier3_price, 2) if tier3_price else None,
            "moq3": moq3,
            "tier4_price": round(tier4_price, 2) if tier4_price else None,
            "gross_weight_kg": round(gross_weight, 2),
            "shipping_template": "Default Shipping Template",
            "lead_time_1": lead1,
            "lead_time_2": lead2,
            "lead_time_3": lead3,
            "pkg_length": 60.0,
            "pkg_width": 42.0,
            "pkg_height": 38.0,
            "status": "Ready for Upload"
        })

    return records

def preview_product_file(filepath=None, file_content=None, filename=""):
    """Reads and returns a lightweight preview of the first records."""
    df, ext, sheet = read_raw_data(filepath, file_content, filename)
    records = parse_records(df.head(5), preserve_images=True)
    return {
        "filename": filename or os.path.basename(filepath or "file"),
        "format": ext.replace('.', '').upper(),
        "sheet": sheet,
        "total_rows": len(df),
        "columns_count": len(df.columns),
        "columns": list(df.columns)[:15],
        "preview_items": records[:3]
    }

def update_database_from_file(filepath=None, file_content=None, filename="", mode="upsert", preserve_images=True):
    """Executes the actual update/import into Database/alibaba_master.db."""
    df, ext, sheet = read_raw_data(filepath, file_content, filename)
    records = parse_records(df, preserve_images=preserve_images)

    conn = get_db()
    cur = conn.cursor()

    if mode == "replace":
        cur.execute("DELETE FROM products")
        conn.commit()

    inserted = 0
    updated = 0

    for p in records:
        # Check if exists
        cur.execute("SELECT id, image_1, image_2, image_3, image_4, image_5, image_6 FROM products WHERE sku = ?", (p['sku'],))
        existing = cur.fetchone()

        # If preserving images and new image is fallback/empty, retain existing images
        img1 = p['image_1']
        img2 = p['image_2']
        img3 = p['image_3']
        img4 = p['image_4']
        img5 = p['image_5']
        img6 = p['image_6']

        if existing and preserve_images and existing['image_1']:
            if not img1 or not img1.startswith('http'):
                img1 = existing['image_1']
                img2 = existing['image_2'] or ''
                img3 = existing['image_3'] or ''
                img4 = existing['image_4'] or ''
                img5 = existing['image_5'] or ''
                img6 = existing['image_6'] or ''

        cur.execute("""
        INSERT INTO products (
            sku, model_number, product_title, category, category_group,
            keywords, image_1, image_2, image_3, image_4, image_5, image_6,
            description_html, place_of_origin, brand_name, outsole_material, upper_material,
            season, feature, gender, sell_by, unit, currency, pricing_type, inventory,
            price_usd, color, size_range, tier1_price, moq1, tier2_price, moq2, tier3_price,
            moq3, tier4_price, gross_weight_kg, shipping_template, lead_time_1, lead_time_2,
            lead_time_3, pkg_length, pkg_width, pkg_height, status
        ) VALUES (
            ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
        ) ON CONFLICT(sku) DO UPDATE SET
            model_number = excluded.model_number,
            product_title = excluded.product_title,
            category = excluded.category,
            category_group = excluded.category_group,
            keywords = excluded.keywords,
            image_1 = ?,
            image_2 = ?,
            image_3 = ?,
            image_4 = ?,
            image_5 = ?,
            image_6 = ?,
            description_html = excluded.description_html,
            place_of_origin = excluded.place_of_origin,
            brand_name = excluded.brand_name,
            outsole_material = excluded.outsole_material,
            upper_material = excluded.upper_material,
            season = excluded.season,
            feature = excluded.feature,
            gender = excluded.gender,
            price_usd = excluded.price_usd,
            color = excluded.color,
            size_range = excluded.size_range,
            tier1_price = excluded.tier1_price,
            moq1 = excluded.moq1,
            tier2_price = excluded.tier2_price,
            moq2 = excluded.moq2,
            tier3_price = excluded.tier3_price,
            moq3 = excluded.moq3,
            gross_weight_kg = excluded.gross_weight_kg,
            status = excluded.status,
            updated_at = CURRENT_TIMESTAMP;
        """, (
            p['sku'], p['model_number'], p['product_title'], p['category'], p['category_group'],
            p['keywords'], img1, img2, img3, img4, img5, img6,
            p['description_html'], p['place_of_origin'], p['brand_name'], p['outsole_material'], p['upper_material'],
            p['season'], p['feature'], p['gender'], p['sell_by'], p['unit'], p['currency'],
            p['pricing_type'], p['inventory'], p['price_usd'], p['color'], p['size_range'],
            p['tier1_price'], p['moq1'], p['tier2_price'], p['moq2'], p['tier3_price'],
            p['moq3'], p['tier4_price'], p['gross_weight_kg'], p['shipping_template'],
            p['lead_time_1'], p['lead_time_2'], p['lead_time_3'], p['pkg_length'], p['pkg_width'],
            p['pkg_height'], p['status'],
            img1, img2, img3, img4, img5, img6
        ))

        if existing:
            updated += 1
        else:
            inserted += 1

    conn.commit()
    conn.close()

    return {
        "success": True,
        "filename": filename or os.path.basename(filepath or "file"),
        "format": ext.replace('.', '').upper(),
        "total": len(records),
        "inserted": inserted,
        "updated": updated,
        "mode": mode
    }

def export_db_to_csv(target_path=EXPORT_PATH_CSV, product_ids=None, regenerate_keywords=False, preferred_angle=None, destination_country=None, update_description=True):
    """Exports all or selected products from database into standard Alibaba CSV format with full 57 columns."""
    conn = get_db()
    cur = conn.cursor()
    if product_ids and len(product_ids) > 0:
        placeholders = ','.join('?' * len(product_ids))
        cur.execute(f"SELECT * FROM products WHERE id IN ({placeholders}) ORDER BY id ASC", list(product_ids))
    else:
        cur.execute("SELECT * FROM products ORDER BY id ASC")
    products = cur.fetchall()
    conn.close()

    with open(target_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        # Row 1: Empty
        writer.writerow([])
        # Row 2: Title notice
        writer.writerow(['Basic Information of Your Listings on Alibaba.com'] + [''] * 56)
        # Row 3: 57 Column Headers
        writer.writerow(ALIBABA_57_COLUMNS)
        # Row 4: Guidelines
        guidelines = [
            '1. Do not enter any unsupported symbols such as !\n2. This field is limited to 128 characters',
            '1. One cell for one image URL.\n2. It contains https:// \n3. The URL must be valid.',
            'Same rule as cell B4', 'Same rule as cell B4', 'Same rule as cell B4', 'Same rule as cell B4',
            'Same rule as cell B4', '1. Product description will be shown on the product page.',
            'Fill them in if useful', 'Fill them in if useful', 'Fill in your preferred category',
            'Attribute name 1', 'Attribute value 1', 'Attribute name 2', 'Attribute value 2',
            'Attribute name 3', 'Attribute value 3', 'Attribute name 4', 'Attribute value 4',
            'Attribute name 5', 'Attribute value 5', 'Select to sell by unit or batch',
            '', 'Unit', 'Currency', 'Pricing type', 'Inventory',
            'SPU ID (Model number)', 'SKU ID', 'SKU price', 'Attribute name 1', 'Attribute value 1',
            'Attribute name 2', 'Attribute value 2', '', '', '', '',
            '1st pricing tier', '1st tier MOQ', '2nd pricing tier', '2nd tier MOQ',
            '3rd pricing tier', '3rd tier MOQ', 'Range price lower', 'Range price higher',
            'Gross weight (KG)', 'Shipping template', 'Shipping qty 1', 'Lead time 1',
            'Shipping qty 2', 'Lead time 2', 'Shipping qty 3', 'Lead time 3',
            'Length', 'Width', 'Height'
        ]
        writer.writerow(guidelines)

        for p in products:
            title_to_export = p['product_title']
            desc_to_export = p['description_html']

            if regenerate_keywords:
                try:
                    try:
                        from accio_keyword_generator import generate_accio_listing
                    except ImportError:
                        from App.accio_keyword_generator import generate_accio_listing

                    pkg = generate_accio_listing(
                        sku=p['sku'],
                        product_title=p['raw_source_title'] or p['product_title'],
                        category_group=p['category_group'] or p['category'],
                        preferred_angle=preferred_angle,
                        model_number=p['model_number'],
                        existing_price=p['price_usd'] or 5.50,
                        existing_weight=p['gross_weight_kg'] or 1.7,
                        destination_country=destination_country
                    )
                    title_to_export = pkg['primary_title']
                    if update_description:
                        desc_to_export = pkg['description_html']
                except Exception as ex:
                    print(f"Accio on-the-fly error for {p['sku']}: {ex}")

            row_data = [
                title_to_export,
                p['image_1'],
                p['image_2'] or '',
                p['image_3'] or '',
                p['image_4'] or '',
                p['image_5'] or '',
                p['image_6'] or '',
                desc_to_export,
                p['place_of_origin'] or 'Indonesia',
                p['brand_name'] or 'AP Boots',
                p['category'] or "Shoes & Accessories > Men's Shoes > Boots",
                "Outsole Material",
                p['outsole_material'] or 'PVC / Rubber',
                "Upper Material",
                p['upper_material'] or 'PVC',
                "Season",
                p['season'] or 'All Seasons',
                "Feature",
                p['feature'] or '100% Waterproof, Anti-Slip',
                "Gender",
                p['gender'] or 'Men / Unisex',
                p['sell_by'] or 'piece',
                '',
                p['unit'] or 'Pair/Pairs',
                p['currency'] or 'USD',
                p['pricing_type'] or 'Tiered pricing',
                p['inventory'] or 5000,
                p['model_number'] or p['sku'],
                p['sku'],
                p['price_usd'] or 5.50,
                "Color",
                p['color'] or 'Black; Green; White; Yellow',
                "Size",
                p['size_range'] or 'EU 38-44',
                '', '', '', '',
                p['tier1_price'] or p['price_usd'],
                p['moq1'] or 100,
                p['tier2_price'] or round((p['price_usd'] or 5.50)*0.95, 2),
                p['moq2'] or 500,
                p['tier3_price'] or round((p['price_usd'] or 5.50)*0.90, 2),
                p['moq3'] or 1000,
                p['tier4_price'] or round((p['price_usd'] or 5.50)*0.85, 2),
                p['tier1_price'] or p['price_usd'],
                p['gross_weight_kg'] or 1.7,
                p['shipping_template'] or 'Default Shipping Template',
                p['moq1'] or 100,
                p['lead_time_1'] or 7,
                p['moq2'] or 500,
                p['lead_time_2'] or 14,
                p['moq3'] or 1000,
                p['lead_time_3'] or 30,
                p['pkg_length'] or 60.0,
                p['pkg_width'] or 42.0,
                p['pkg_height'] or 38.0
            ]
            writer.writerow(row_data)

    return len(products)

def export_db_to_excel(target_path=EXPORT_PATH_XLSX, product_ids=None, regenerate_keywords=False, preferred_angle=None, destination_country=None, update_description=True):
    """Exports all or selected products into standard 57-column Alibaba Excel workbook."""
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb['Alibaba default template']
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

    conn = get_db()
    cur = conn.cursor()
    if product_ids and len(product_ids) > 0:
        placeholders = ','.join('?' * len(product_ids))
        cur.execute(f"SELECT * FROM products WHERE id IN ({placeholders}) ORDER BY id ASC", list(product_ids))
    else:
        cur.execute("SELECT * FROM products ORDER BY id ASC")
    products = cur.fetchall()
    conn.close()

    for idx, p in enumerate(products):
        row_idx = 5 + idx
        title_to_export = p['product_title']
        desc_to_export = p['description_html']

        if regenerate_keywords:
            try:
                try:
                    from accio_keyword_generator import generate_accio_listing
                except ImportError:
                    from App.accio_keyword_generator import generate_accio_listing
                pkg = generate_accio_listing(
                    sku=p['sku'],
                    product_title=p['raw_source_title'] or p['product_title'],
                    category_group=p['category_group'] or p['category'],
                    preferred_angle=preferred_angle,
                    model_number=p['model_number'],
                    existing_price=p['price_usd'] or 5.50,
                    existing_weight=p['gross_weight_kg'] or 1.7,
                    destination_country=destination_country
                )
                title_to_export = pkg['primary_title']
                if update_description:
                    desc_to_export = pkg['description_html']
            except Exception as ex:
                print(f"Accio on-the-fly error for {p['sku']}: {ex}")

        row_data = [
            title_to_export,
            p['image_1'],
            p['image_2'] or '',
            p['image_3'] or '',
            p['image_4'] or '',
            p['image_5'] or '',
            p['image_6'] or '',
            desc_to_export,
            p['place_of_origin'] or 'Indonesia',
            p['brand_name'] or 'AP Boots',
            p['category'] or "Shoes & Accessories > Men's Shoes > Boots",
            "Outsole Material",
            p['outsole_material'] or 'PVC / Rubber',
            "Upper Material",
            p['upper_material'] or 'PVC',
            "Season",
            p['season'] or 'All Seasons',
            "Feature",
            p['feature'] or '100% Waterproof, Anti-Slip',
            "Gender",
            p['gender'] or 'Men / Unisex',
            p['sell_by'] or 'piece',
            None,
            p['unit'] or 'Pair/Pairs',
            p['currency'] or 'USD',
            p['pricing_type'] or 'Tiered pricing',
            p['inventory'] or 5000,
            p['model_number'] or p['sku'],
            p['sku'],
            p['price_usd'] or 5.50,
            "Color",
            p['color'] or 'Black; Green; White; Yellow',
            "Size",
            p['size_range'] or 'EU 38-44',
            None, None, None, None,
            p['tier1_price'] or p['price_usd'],
            p['moq1'] or 100,
            p['tier2_price'] or round((p['price_usd'] or 5.50)*0.95, 2),
            p['moq2'] or 500,
            p['tier3_price'] or round((p['price_usd'] or 5.50)*0.90, 2),
            p['moq3'] or 1000,
            p['tier4_price'] or round((p['price_usd'] or 5.50)*0.85, 2),
            p['tier1_price'] or p['price_usd'],
            p['gross_weight_kg'] or 1.7,
            p['shipping_template'] or 'Default Shipping Template',
            p['moq1'] or 100,
            p['lead_time_1'] or 7,
            p['moq2'] or 500,
            p['lead_time_2'] or 14,
            p['moq3'] or 1000,
            p['lead_time_3'] or 30,
            p['pkg_length'] or 60.0,
            p['pkg_width'] or 42.0,
            p['pkg_height'] or 38.0
        ]
        for c_i, val in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=c_i, value=val)

    wb.save(target_path)
    return len(products)

def cli_main():
    import sys
    print("===================================================================")
    print("   ALIBABA LISTING PRODUCT DATA LOADER & UPDATER (XLSX / CSV)")
    print("             PT Indo Sukses Semesta - AP Boots")
    print("===================================================================")

    if len(sys.argv) > 1:
        arg = sys.argv[1].strip()
        if arg in ['--export-csv', '-ec']:
            cnt = export_db_to_csv()
            print(f"[OK] Berhasil mengekspor {cnt} produk ke {EXPORT_PATH_CSV}")
            return
        elif arg in ['--export-xlsx', '-ex']:
            cnt = export_db_to_excel()
            print(f"[OK] Berhasil mengekspor {cnt} produk ke {EXPORT_PATH_XLSX}")
            return
        elif os.path.exists(arg):
            target_path = arg
        else:
            target_path = os.path.join(BASE_DIR, arg)
            if not os.path.exists(target_path):
                print(f"[ERROR] File tidak ditemukan: {arg}")
                return
    else:
        # Interactive selection
        files = list_available_files()
        if not files:
            print("[INFO] Tidak ada file .xlsx atau .csv ditemukan di Input/ atau Output/")
            return

        print("\nPilih file yang ingin diimpor / diupdate ke database:")
        for idx, f in enumerate(files, start=1):
            print(f" [{idx:2d}] [{f['format']:4s}] {f['id']:40s} ({f['size_kb']} KB, {f['modified']})")

        print(" [ 0] Batal")
        try:
            choice = input("\nMasukkan nomor pilihan (0-" + str(len(files)) + "): ").strip()
            if not choice or choice == '0':
                print("[INFO] Operasi dibatalkan.")
                return
            idx_choice = int(choice) - 1
            if idx_choice < 0 or idx_choice >= len(files):
                print("[ERROR] Pilihan tidak valid.")
                return
            target_path = files[idx_choice]['path']
        except (ValueError, KeyboardInterrupt):
            print("\n[INFO] Batal.")
            return

    print(f"\n[1/3] Membaca dan menganalisis file: {target_path}")
    prev = preview_product_file(filepath=target_path)
    print(f"      Format: {prev['format']}, Total Baris: {prev['total_rows']}, Kolom: {prev['columns_count']}")

    print("\n[2/3] Memperbarui database alibaba_master.db (Mode: Upsert)...")
    res = update_database_from_file(filepath=target_path, mode='upsert', preserve_images=True)

    print(f"\n[3/3] SELESAI!")
    print(f"      - Produk Baru Dimasukkan : {res['inserted']}")
    print(f"      - Produk Diperbarui      : {res['updated']}")
    print(f"      - Total Record Diproses  : {res['total']}")
    print("===================================================================")

if __name__ == '__main__':
    cli_main()

