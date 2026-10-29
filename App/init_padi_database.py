import openpyxl
import sqlite3
import glob
import os
import re

BASE_DIR = r"d:\Work\Hots\Product Input"
DB_PATH = os.path.join(BASE_DIR, "Database", "padi_umkm_master.db") if os.path.exists(os.path.join(BASE_DIR, "Database", "padi_umkm_master.db")) else os.path.join(BASE_DIR, "padi_umkm_master.db")
OUTPUT_PADI = os.path.join(BASE_DIR, "Output_PadiUMKM")

def init_db():
    print(f"Connecting to database: {DB_PATH}")
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    cur.execute("""
    CREATE TABLE IF NOT EXISTS padi_products (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        batch_num INTEGER,
        row_in_batch INTEGER,
        batch_filename TEXT,
        tipe_produk TEXT DEFAULT 'Barang',
        nama_produk TEXT NOT NULL,
        deskripsi_produk TEXT NOT NULL,
        kategori_kode INTEGER,
        berat_produk_gr INTEGER,
        dimensi_panjang_cm INTEGER,
        dimensi_lebar_cm INTEGER,
        dimensi_tinggi_cm INTEGER,
        pembelian_minimum INTEGER DEFAULT 1,
        brand TEXT DEFAULT 'AP Boots',
        pdn TEXT DEFAULT 'Ya',
        jenis_tkdn TEXT DEFAULT 'Tidak Mengetahui',
        nilai_tkdn REAL,
        dokumen_pendukung_tkdn TEXT,
        foto_produk_1 TEXT,
        foto_produk_2 TEXT,
        foto_produk_3 TEXT,
        foto_produk_4 TEXT,
        foto_produk_5 TEXT,
        url_video_produk_1 TEXT,
        stok_unit INTEGER DEFAULT 100,
        satuan_stok TEXT DEFAULT 'pcs',
        nama_satuan_lainnya TEXT,
        stok_selalu_ada TEXT DEFAULT 'Aktif',
        pre_order_duration INTEGER,
        harga_satuan INTEGER,
        ppn TEXT DEFAULT 'Non Aktif',
        opsi_pph TEXT DEFAULT 'PPH 22',
        asuransi_pengiriman TEXT DEFAULT 'Wajib',
        jadikan_gratis_ongkir TEXT DEFAULT 'Tidak',
        visibilitas_produk TEXT DEFAULT 'Publik',
        kode_sku TEXT,
        status TEXT DEFAULT 'Ready',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    """)
    conn.commit()
    return conn

def import_batches_to_db():
    conn = init_db()
    cur = conn.cursor()
    
    # Reset table for fresh, clean import
    cur.execute("DELETE FROM padi_products")
    cur.execute("DELETE FROM sqlite_sequence WHERE name='padi_products'")
    conn.commit()

    batch_files = sorted(glob.glob(os.path.join(OUTPUT_PADI, "PadiUMKM_Bulk_Upload_Batch_*.xlsx")))
    if not batch_files:
        print("[ERROR] No batch files found in Output_PadiUMKM!")
        return

    total_inserted = 0

    for b_idx, bf in enumerate(batch_files, 1):
        filename = os.path.basename(bf)
        print(f"\nProcessing Batch {b_idx}: {filename}")
        wb = openpyxl.load_workbook(bf, data_only=True)
        ws = wb['Isian Template Upload Barang']

        for r in range(4, ws.max_row + 1):
            nama_produk = ws.cell(r, 3).value
            if not nama_produk:
                continue

            tipe_produk = ws.cell(r, 2).value or 'Barang'
            deskripsi = ws.cell(r, 4).value or ''
            kat_kode = ws.cell(r, 5).value or 2450
            berat = ws.cell(r, 6).value or 1000
            pjg = ws.cell(r, 7).value or 30
            lbr = ws.cell(r, 8).value or 20
            tgi = ws.cell(r, 9).value or 15
            min_order = ws.cell(r, 10).value or 1
            brand = ws.cell(r, 11).value or 'AP Boots'
            pdn = ws.cell(r, 12).value or 'Ya'
            jenis_tkdn = ws.cell(r, 13).value or 'Tidak Mengetahui'
            nilai_tkdn = ws.cell(r, 14).value
            doc_tkdn = ws.cell(r, 15).value
            foto1 = ws.cell(r, 16).value or ''
            foto2 = ws.cell(r, 17).value or ''
            foto3 = ws.cell(r, 18).value or ''
            foto4 = ws.cell(r, 19).value or ''
            foto5 = ws.cell(r, 20).value or ''

            # Shopee: non-main image jangan di-load (hanya foto1 / main image yang tetap)
            if 'susercontent.com' in str(foto2): foto2 = ''
            if 'susercontent.com' in str(foto3): foto3 = ''
            if 'susercontent.com' in str(foto4): foto4 = ''
            if 'susercontent.com' in str(foto5): foto5 = ''
            
            # Check for column shift (where Col 22 has 'pcs' instead of integer)
            col21_val = ws.cell(r, 21).value
            col22_val = ws.cell(r, 22).value
            
            if col22_val == 'pcs' or col21_val == 100 or col21_val == '100':
                # Shifted row detected!
                video = ''
                stok = 100
                satuan = 'pcs'
                satuan_lain = None
                stok_selalu_ada = 'Aktif'
                preorder = None
                harga = ws.cell(r, 26).value or 0
                ppn = ws.cell(r, 27).value or 'Non Aktif'
                pph = ws.cell(r, 28).value or 'PPH 22'
                asuransi = ws.cell(r, 29).value or 'Wajib'
                gratis_ongkir = ws.cell(r, 30).value or 'Tidak'
                visibilitas = ws.cell(r, 31).value or 'Publik'
                sku = ws.cell(r, 33).value or f"APB-BATCH{b_idx}-{r}"
            else:
                video = ws.cell(r, 21).value or ''
                stok = ws.cell(r, 22).value or 100
                satuan = ws.cell(r, 23).value or 'pcs'
                satuan_lain = ws.cell(r, 24).value
                stok_selalu_ada = ws.cell(r, 25).value or 'Aktif'
                preorder = ws.cell(r, 26).value
                harga = ws.cell(r, 27).value or 0
                ppn = ws.cell(r, 28).value or 'Non Aktif'
                pph = ws.cell(r, 29).value or 'PPH 22'
                asuransi = ws.cell(r, 30).value or 'Wajib'
                gratis_ongkir = ws.cell(r, 31).value or 'Tidak'
                visibilitas = ws.cell(r, 32).value or 'Publik'
                sku = ws.cell(r, 33).value or f"APB-BATCH{b_idx}-{r}"

            # Clean price if string
            if isinstance(harga, str):
                harga_cleaned = re.sub(r'[^\d]', '', harga)
                harga = int(harga_cleaned) if harga_cleaned else 0

            cur.execute("""
            INSERT INTO padi_products (
                batch_num, row_in_batch, batch_filename, tipe_produk, nama_produk, deskripsi_produk,
                kategori_kode, berat_produk_gr, dimensi_panjang_cm, dimensi_lebar_cm, dimensi_tinggi_cm,
                pembelian_minimum, brand, pdn, jenis_tkdn, nilai_tkdn, dokumen_pendukung_tkdn,
                foto_produk_1, foto_produk_2, foto_produk_3, foto_produk_4, foto_produk_5,
                url_video_produk_1, stok_unit, satuan_stok, nama_satuan_lainnya, stok_selalu_ada,
                pre_order_duration, harga_satuan, ppn, opsi_pph, asuransi_pengiriman,
                jadikan_gratis_ongkir, visibilitas_produk, kode_sku
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                b_idx, r, filename, str(tipe_produk), str(nama_produk), str(deskripsi),
                int(kat_kode) if kat_kode else 2450, int(berat) if berat else 1000,
                int(pjg) if pjg else 30, int(lbr) if lbr else 20, int(tgi) if tgi else 15,
                int(min_order) if min_order else 1, str(brand), str(pdn),
                str(jenis_tkdn), float(nilai_tkdn) if nilai_tkdn else None, str(doc_tkdn) if doc_tkdn else None,
                str(foto1), str(foto2), str(foto3), str(foto4), str(foto5),
                str(video) if video else None, int(stok) if stok else 100, str(satuan),
                str(satuan_lain) if satuan_lain else None, str(stok_selalu_ada),
                int(preorder) if preorder else None, int(harga) if harga else 0,
                str(ppn), str(pph), str(asuransi), str(gratis_ongkir), str(visibilitas), str(sku)
            ))
            total_inserted += 1

    conn.commit()
    conn.close()
    print(f"\n[SUCCESS] Successfully imported {total_inserted} products into SQLite '{DB_PATH}'.")

if __name__ == '__main__':
    import_batches_to_db()
