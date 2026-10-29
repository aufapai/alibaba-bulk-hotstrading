import openpyxl
import sqlite3
import os
import re

BASE_DIR = r"d:\Work\Hots\Product Input"
DB_PATH = os.path.join(BASE_DIR, "Database", "alibaba_master.db") if os.path.exists(os.path.join(BASE_DIR, "Database", "alibaba_master.db")) else os.path.join(BASE_DIR, "alibaba_master.db")
EXCEL_PATH = os.path.join(BASE_DIR, "Output", "Alibaba_Bulk_Upload_Tokopedia_Master_2026.xlsx") if os.path.exists(os.path.join(BASE_DIR, "Output", "Alibaba_Bulk_Upload_Tokopedia_Master_2026.xlsx")) else os.path.join(BASE_DIR, "Alibaba_Bulk_Upload_Tokopedia_Master_2026.xlsx")
SUCCESS_PATH = os.path.join(BASE_DIR, "Output", "id29233741662kjoc_368505_success.xlsx") if os.path.exists(os.path.join(BASE_DIR, "Output", "id29233741662kjoc_368505_success.xlsx")) else os.path.join(BASE_DIR, "id29233741662kjoc_368505_success.xlsx")

def deduplicate_images(img_list):
    seen = set()
    unique_imgs = []
    for img in img_list:
        if not img or not isinstance(img, str):
            continue
        clean_img = img.strip()
        if not clean_img.startswith("http"):
            continue
        # Normalize for deduplication comparison (ignore query string if any)
        norm_key = clean_img.split("?")[0].rstrip("/")
        if norm_key not in seen:
            seen.add(norm_key)
            unique_imgs.append(clean_img)
            
    # Pad with empty strings up to 6
    while len(unique_imgs) < 6:
        unique_imgs.append("")
    return unique_imgs[:6]

def generate_seo_keywords(title, category):
    kw_list = []
    t_low = title.lower()
    if "terra" in t_low:
        kw_list = ["pvc rain boots", "waterproof work boots", "agriculture safety boots"]
    elif "orca" in t_low:
        kw_list = ["industrial rain boots", "high waterproof boots", "rubber safety shoes"]
    elif "chef" in t_low:
        kw_list = ["kitchen work shoes", "slip resistant chef shoes", "waterproof slip on"]
    elif "medic" in t_low:
        kw_list = ["medical safety shoes", "cleanroom protective footwear", "anti slip clogs"]
    elif "allbike" in t_low or "moto" in t_low:
        kw_list = ["motorcycle riding shoes", "waterproof biker shoes", "breathable rain boots"]
    elif "safety" in t_low or "s5" in t_low:
        kw_list = ["steel toe safety boots", "s5 industrial work boots", "chemical resistant boots"]
    elif "star" in t_low:
        kw_list = ["casual waterproof sneakers", "pvc canvas style shoes", "daily rain kets"]
    elif "kids" in t_low:
        kw_list = ["kids rain boots", "waterproof children boots", "cute toddler rubber boots"]
    elif "clogs" in t_low or "995" in t_low or "sandal" in t_low:
        kw_list = ["eva pvc clogs", "anti slip work slippers", "waterproof garden shoes"]
    else:
        kw_list = ["waterproof pvc boots", "heavy duty rain boots", "indonesian safety footwear"]
    return ", ".join(kw_list)

def init_db():
    print(f"Connecting to database: {DB_PATH}")
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    cur.execute("""
    CREATE TABLE IF NOT EXISTS products (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sku TEXT UNIQUE NOT NULL,
        model_number TEXT,
        alibaba_product_id TEXT,
        product_title TEXT NOT NULL,
        category TEXT NOT NULL,
        category_group TEXT,
        keywords TEXT,
        image_1 TEXT NOT NULL,
        image_2 TEXT,
        image_3 TEXT,
        image_4 TEXT,
        image_5 TEXT,
        image_6 TEXT,
        description_html TEXT NOT NULL,
        place_of_origin TEXT DEFAULT 'Indonesia',
        brand_name TEXT DEFAULT 'AP Boots',
        outsole_material TEXT DEFAULT 'PVC / Rubber',
        upper_material TEXT DEFAULT 'PVC',
        season TEXT DEFAULT 'All Seasons',
        feature TEXT,
        gender TEXT DEFAULT 'Men / Unisex',
        sell_by TEXT DEFAULT 'piece',
        unit TEXT DEFAULT 'Pair/Pairs',
        currency TEXT DEFAULT 'USD',
        pricing_type TEXT DEFAULT 'Tiered pricing',
        inventory INTEGER DEFAULT 5000,
        price_usd REAL DEFAULT 5.50,
        color TEXT DEFAULT 'Black; Green; White; Yellow',
        size_range TEXT DEFAULT 'EU 38-44',
        tier1_price REAL,
        moq1 INTEGER DEFAULT 100,
        tier2_price REAL,
        moq2 INTEGER DEFAULT 500,
        tier3_price REAL,
        moq3 INTEGER DEFAULT 1000,
        tier4_price REAL,
        gross_weight_kg REAL DEFAULT 1.7,
        shipping_template TEXT DEFAULT 'Default Shipping Template',
        lead_time_1 INTEGER DEFAULT 7,
        lead_time_2 INTEGER DEFAULT 14,
        lead_time_3 INTEGER DEFAULT 30,
        pkg_length REAL DEFAULT 60.0,
        pkg_width REAL DEFAULT 42.0,
        pkg_height REAL DEFAULT 38.0,
        status TEXT DEFAULT 'Ready',
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)
    conn.commit()

    # Load Success Report from Alibaba if present
    success_map = {}
    if os.path.exists(SUCCESS_PATH):
        print(f"Loading Alibaba success report from: {SUCCESS_PATH}")
        swb = openpyxl.load_workbook(SUCCESS_PATH, data_only=True)
        sws = swb['success'] if 'success' in swb.sheetnames else swb.active
        for r in list(sws.iter_rows(values_only=True))[1:]:
            model_num = str(r[0]).strip() if r[0] else ""
            p_id = str(r[1]).strip() if len(r) > 1 and r[1] else ""
            if model_num and p_id:
                success_map[model_num] = p_id
        print(f"Found {len(success_map)} verified live listings on Alibaba!")

    # Read from Alibaba Master Template Excel
    print(f"Reading master Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Alibaba default template']
    rows = list(ws.iter_rows(values_only=True))

    inserted = 0
    updated = 0

    for r_idx, r in enumerate(rows[4:], start=5):
        if not r or not r[0]:
            continue

        title = str(r[0]).strip()
        raw_imgs = [r[1], r[2], r[3], r[4], r[5], r[6]]
        # Shopee: non-main image jangan di-load (hanya main image raw_imgs[0] yang boleh)
        for i in range(1, len(raw_imgs)):
            if raw_imgs[i] and 'susercontent.com' in str(raw_imgs[i]):
                raw_imgs[i] = None
        dedup_imgs = deduplicate_images(raw_imgs)

        desc = str(r[7]).strip() if r[7] else ""
        origin = str(r[8]).strip() if r[8] else "Indonesia"
        brand = str(r[9]).strip() if r[9] else "AP Boots"
        cat = str(r[10]).strip() if r[10] else "Shoes & Accessories > Men's Shoes > Boots"
        outsole = str(r[12]).strip() if len(r) > 12 and r[12] else "PVC / Rubber"
        upper = str(r[14]).strip() if len(r) > 14 and r[14] else "PVC"
        season = str(r[16]).strip() if len(r) > 16 and r[16] else "All Seasons"
        feature = str(r[18]).strip() if len(r) > 18 and r[18] else "100% Waterproof, Anti-Slip"
        gender = str(r[20]).strip() if len(r) > 20 and r[20] else "Men / Unisex"
        sell_by = str(r[21]).strip() if len(r) > 21 and r[21] else "piece"
        unit = str(r[23]).strip() if len(r) > 23 and r[23] else "Pair/Pairs"
        currency = str(r[24]).strip() if len(r) > 24 and r[24] else "USD"
        pricing_type = str(r[25]).strip() if len(r) > 25 and r[25] else "Tiered pricing"
        inventory = int(r[26]) if len(r) > 26 and r[26] else 5000
        model = str(r[27]).strip() if len(r) > 27 and r[27] else f"APB-{r_idx-4:03d}"
        sku = str(r[28]).strip() if len(r) > 28 and r[28] else model
        price_usd = float(r[29]) if len(r) > 29 and r[29] else 5.50
        color = str(r[31]).strip() if len(r) > 31 and r[31] else "Black; Green; White; Yellow"
        size_range = str(r[33]).strip() if len(r) > 33 and r[33] else "EU 38-44"

        t1_price = float(r[38]) if len(r) > 38 and r[38] else price_usd
        moq1 = int(r[39]) if len(r) > 39 and r[39] else 100
        t2_price = float(r[40]) if len(r) > 40 and r[40] else round(t1_price * 0.95, 2)
        moq2 = int(r[41]) if len(r) > 41 and r[41] else 500
        t3_price = float(r[42]) if len(r) > 42 and r[42] else round(t1_price * 0.90, 2)
        moq3 = int(r[43]) if len(r) > 43 and r[43] else 1000
        t4_price = float(r[44]) if len(r) > 44 and r[44] else round(t1_price * 0.85, 2)

        gross_wt = float(r[46]) if len(r) > 46 and r[46] else 1.7
        ship_tmpl = str(r[47]).strip() if len(r) > 47 and r[47] else "Default Shipping Template"
        lead1 = int(r[49]) if len(r) > 49 and r[49] else 7
        lead2 = int(r[51]) if len(r) > 51 and r[51] else 14
        lead3 = int(r[53]) if len(r) > 53 and r[53] else 30
        pkg_l = float(r[54]) if len(r) > 54 and r[54] else 60.0
        pkg_w = float(r[55]) if len(r) > 55 and r[55] else 42.0
        pkg_h = float(r[56]) if len(r) > 56 and r[56] else 38.0

        # Alibaba Product ID from success map
        ali_id = success_map.get(model) or success_map.get(sku)
        status = f"Live on Alibaba (ID: {ali_id})" if ali_id else "Ready for Upload"

        # Determine category group
        t_low = title.lower()
        if "terra" in t_low: cat_grp = "AP Terra Series"
        elif "orca" in t_low: cat_grp = "AP Orca Series"
        elif "chef" in t_low: cat_grp = "AP Chef Series"
        elif "medic" in t_low: cat_grp = "AP Medic Series"
        elif "safety" in t_low or "s5" in t_low: cat_grp = "AP Safety Boots"
        elif "star" in t_low: cat_grp = "AP Kasual Sneakers"
        elif "allbike" in t_low or "moto" in t_low or "hobby" in t_low: cat_grp = "AP Biker & Casual"
        elif "kids" in t_low: cat_grp = "AP Kids Boots"
        elif "clogs" in t_low or "sandal" in t_low: cat_grp = "AP Sandals & Clogs"
        else: cat_grp = "AP Classic Boots"

        keywords = generate_seo_keywords(title, cat)

        cur.execute("""
        INSERT INTO products (
            sku, model_number, alibaba_product_id, product_title, category, category_group,
            keywords, image_1, image_2, image_3, image_4, image_5, image_6,
            description_html, place_of_origin, brand_name, outsole_material, upper_material,
            season, feature, gender, sell_by, unit, currency, pricing_type, inventory,
            price_usd, color, size_range, tier1_price, moq1, tier2_price, moq2, tier3_price,
            moq3, tier4_price, gross_weight_kg, shipping_template, lead_time_1, lead_time_2,
            lead_time_3, pkg_length, pkg_width, pkg_height, status
        ) VALUES (
            ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
        ) ON CONFLICT(sku) DO UPDATE SET
            alibaba_product_id = excluded.alibaba_product_id,
            product_title = excluded.product_title,
            keywords = excluded.keywords,
            image_1 = excluded.image_1,
            image_2 = excluded.image_2,
            image_3 = excluded.image_3,
            image_4 = excluded.image_4,
            image_5 = excluded.image_5,
            image_6 = excluded.image_6,
            description_html = excluded.description_html,
            price_usd = excluded.price_usd,
            tier1_price = excluded.tier1_price,
            gross_weight_kg = excluded.gross_weight_kg,
            status = excluded.status,
            updated_at = CURRENT_TIMESTAMP;
        """, (
            sku, model, ali_id, title, cat, cat_grp, keywords,
            dedup_imgs[0], dedup_imgs[1], dedup_imgs[2], dedup_imgs[3], dedup_imgs[4], dedup_imgs[5],
            desc, origin, brand, outsole, upper, season, feature, gender, sell_by, unit, currency,
            pricing_type, inventory, price_usd, color, size_range, t1_price, moq1, t2_price, moq2,
            t3_price, moq3, t4_price, gross_wt, ship_tmpl, lead1, lead2, lead3, pkg_l, pkg_w, pkg_h, status
        ))
        inserted += 1

    conn.commit()
    conn.close()
    print(f"Database initialized successfully! Processed {inserted} products.")

if __name__ == '__main__':
    init_db()
