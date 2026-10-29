import openpyxl
import re
import os
import glob
from collections import Counter

# 1. SHOPEE ANALYSIS
swb = openpyxl.load_workbook('Input/shopee-co-id-2026-09-23.xlsx', data_only=True)
srows = list(swb.active.iter_rows(values_only=True))
shopee_items = []
for r in srows[1:]:
    name = r[4]
    price_raw = r[3]
    sold_raw = r[5]
    rating_raw = r[7]
    badge = r[9]
    # Shopee: hanya main image r[11] yang di-load; non-main image jangan di-load
    imgs = [r[11]] if r[11] else []
    
    p_clean = None
    if price_raw:
        m = re.search(r'([0-9\.]+)', str(price_raw))
        if m:
            p_clean = int(m.group(1).replace('.', ''))
            
    s_clean = 0
    if sold_raw:
        txt = str(sold_raw).lower().replace('terjual', '').strip()
        if 'rb' in txt:
            txt = txt.replace('rb', '').replace('+', '').strip().replace(',', '.')
            try:
                s_clean = int(float(txt) * 1000)
            except:
                s_clean = 0
        else:
            try:
                s_clean = int(re.sub(r'[^0-9]', '', txt))
            except:
                s_clean = 0
                
    shopee_items.append({
        'name': name,
        'price': p_clean,
        'sold': s_clean,
        'sold_raw': sold_raw,
        'rating': float(rating_raw) if rating_raw else None,
        'badge': badge,
        'images_count': len(imgs)
    })

prices = [x['price'] for x in shopee_items if x['price']]
solds = [x['sold'] for x in shopee_items if x['sold']]
ratings = [x['rating'] for x in shopee_items if x['rating']]

print("=== SHOPEE SUMMARY ===")
print(f"Total Rows: {len(shopee_items)}")
print(f"Price: Min Rp{min(prices):,}, Max Rp{max(prices):,}, Avg Rp{int(sum(prices)/len(prices)):,}")
print(f"Total Recorded Sold (est): {sum(solds):,} pcs")
print(f"Rating Avg: {sum(ratings)/len(ratings):.2f} (from {len(ratings)} products)")
print("\nTop 10 Shopee Best Sellers:")
shopee_sorted = sorted(shopee_items, key=lambda x: x['sold'], reverse=True)
for i, it in enumerate(shopee_sorted[:10], 1):
    print(f" {i}. {it['name']} | Sold: {it['sold_raw']} | Price: Rp{it['price']:,} | Rating: {it['rating']}")

# 2. TOKOPEDIA DETAILED ANALYSIS
tokopedia_dict = {}

tokopedia_files = [
    ('Input/tokopedia-com-2026-09-23.xlsx', 'Batch 1 (20 items)'),
    ('Input/tokopedia-com-2026-09-23-3.xlsx', 'Batch 2 (81 items)'),
    ('Input/tokopedia-com-2026-09-23-3 (1).xlsx', 'Batch 3 (114 items)')
]

for filepath, label in tokopedia_files:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    col_map = {h: i for i, h in enumerate(header)}
    
    print(f"\nProcessing {label} -> {filepath}: {len(rows)-1} rows")
    
    for r in rows[1:]:
        link = r[col_map.get('item_page_link')] if col_map.get('item_page_link') is not None else None
        if not link:
            continue
        clean_url = link.split('?')[0].rstrip('/')
        
        name = r[col_map['product_name']] if col_map.get('product_name') is not None else None
        list_name = r[col_map['data']] if col_map.get('data') is not None else None
        price_raw = r[col_map['price']] if col_map.get('price') is not None else None
        orig_price = r[col_map['original_price']] if col_map.get('original_price') is not None else None
        discount = r[col_map['discount']] if col_map.get('discount') is not None else None
        rating_raw = r[col_map['rating']] if col_map.get('rating') is not None else None
        rev_count = r[col_map['review_count']] if col_map.get('review_count') is not None else None
        sold_raw = r[col_map['sold']] if col_map.get('sold') is not None else None
        desc = r[col_map['description']] if col_map.get('description') is not None else None
        stock = r[col_map['stock_level']] if col_map.get('stock_level') is not None else None
        main_img = r[col_map['main_image']] if col_map.get('main_image') is not None else None
        gallery = r[col_map['image_17']] if col_map.get('image_17') is not None else None
        
        if clean_url not in tokopedia_dict:
            tokopedia_dict[clean_url] = {
                'url': clean_url,
                'name': name or list_name,
                'price_raw': price_raw,
                'orig_price': orig_price,
                'discount': discount,
                'rating_raw': rating_raw,
                'review_count': rev_count,
                'sold_raw': sold_raw,
                'desc': desc,
                'stock': stock,
                'main_img': main_img,
                'gallery': gallery,
                'batches': [label]
            }
        else:
            tokopedia_dict[clean_url]['batches'].append(label)
            # enrich if empty
            if not tokopedia_dict[clean_url]['desc'] and desc:
                tokopedia_dict[clean_url]['desc'] = desc
            if not tokopedia_dict[clean_url]['name'] and (name or list_name):
                tokopedia_dict[clean_url]['name'] = name or list_name
            if not tokopedia_dict[clean_url]['price_raw'] and price_raw:
                tokopedia_dict[clean_url]['price_raw'] = price_raw
            if not tokopedia_dict[clean_url]['sold_raw'] and sold_raw:
                tokopedia_dict[clean_url]['sold_raw'] = sold_raw
            if not tokopedia_dict[clean_url]['stock'] and stock:
                tokopedia_dict[clean_url]['stock'] = stock
            if not tokopedia_dict[clean_url]['main_img'] and main_img:
                tokopedia_dict[clean_url]['main_img'] = main_img
            if not tokopedia_dict[clean_url]['gallery'] and gallery:
                tokopedia_dict[clean_url]['gallery'] = gallery
            if not tokopedia_dict[clean_url]['orig_price'] and orig_price:
                tokopedia_dict[clean_url]['orig_price'] = orig_price
            if not tokopedia_dict[clean_url]['discount'] and discount:
                tokopedia_dict[clean_url]['discount'] = discount

print(f"\n=== COMBINED TOKOPEDIA METRICS ===")
print(f"Total Unique Product URLs: {len(tokopedia_dict)}")
desc_count = sum(1 for v in tokopedia_dict.values() if v['desc'])
print(f"Products with Full Description: {desc_count}")
print(f"Products with Image: {sum(1 for v in tokopedia_dict.values() if v['main_img'])}")

# Analyze Tokopedia descriptions: Extract Categories, Etalase, Weights, Sizes
etalase_counter = Counter()
kategori_counter = Counter()
berat_list = []
tokopedia_prices = []
tokopedia_solds = []

for v in tokopedia_dict.values():
    # Price
    if v['price_raw']:
        m = re.search(r'([0-9\.]+)', str(v['price_raw']))
        if m:
            try:
                p = int(m.group(1).replace('.', ''))
                tokopedia_prices.append(p)
                v['price_clean'] = p
            except:
                pass
                
    # Sold
    if v['sold_raw']:
        txt = str(v['sold_raw']).lower().replace('terjual', '').strip()
        s_val = 0
        if 'rb' in txt:
            txt = txt.replace('rb', '').replace('+', '').strip().replace(',', '.')
            try:
                s_val = int(float(txt) * 1000)
            except:
                pass
        else:
            try:
                s_val = int(re.sub(r'[^0-9]', '', txt))
            except:
                pass
        v['sold_clean'] = s_val
        tokopedia_solds.append(s_val)
    else:
        v['sold_clean'] = 0

    if v['desc']:
        # Extract Etalase: ...
        m_et = re.search(r'Etalase:\s*([^\n\r]+)', v['desc'])
        if m_et:
            et = m_et.group(1).strip()
            etalase_counter[et] += 1
            v['etalase'] = et
            
        # Extract Kategori: ...
        m_kat = re.search(r'Kategori:\s*([^\n\r]+)', v['desc'])
        if m_kat:
            kat = m_kat.group(1).strip()
            kategori_counter[kat] += 1
            v['kategori'] = kat
            
        # Extract Berat Satuan: ...
        m_brt = re.search(r'Berat Satuan:\s*([^\n\r]+)', v['desc'])
        if m_brt:
            berat_list.append(m_brt.group(1).strip())

print("\n--- Tokopedia Etalase Breakdown ---")
for et, count in etalase_counter.most_common():
    print(f"  {et}: {count} products")

print("\n--- Tokopedia Kategori Breakdown ---")
for kat, count in kategori_counter.most_common():
    print(f"  {kat}: {count} products")

if tokopedia_prices:
    print(f"\nTokopedia Price Stats: Min Rp{min(tokopedia_prices):,}, Max Rp{max(tokopedia_prices):,}, Avg Rp{int(sum(tokopedia_prices)/len(tokopedia_prices)):,}")

print("\nTop 10 Tokopedia Best Sellers:")
tok_sorted = sorted(tokopedia_dict.values(), key=lambda x: x.get('sold_clean', 0), reverse=True)
for i, it in enumerate(tok_sorted[:10], 1):
    print(f" {i}. {it['name']} | Sold: {it['sold_raw']} | Price: {it['price_raw']} | Rating: {it['rating_raw']}")

