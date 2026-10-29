import os
import sqlite3
import openpyxl
from datetime import datetime
import re

BASE_DIR = r"d:\Work\Hots\Product Input"
DB_PATH = os.path.join(BASE_DIR, "Database", "alibaba_master.db")
INPUT_DIR = os.path.join(BASE_DIR, "Input")
OUTPUT_DIR = os.path.join(BASE_DIR, "Output")
TEMPLATE_PATH = os.path.join(BASE_DIR, "Template", "Alibaba default template (recommended) (1).xlsx")

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_upload_reports_table():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
    CREATE TABLE IF NOT EXISTS alibaba_upload_reports (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        file_name TEXT NOT NULL,
        file_row INTEGER,
        sku TEXT NOT NULL,
        model_number TEXT,
        product_title TEXT,
        failed_reason TEXT NOT NULL,
        error_category TEXT,
        original_img1 TEXT,
        original_img2 TEXT,
        resolved_img1 TEXT,
        resolved_img2 TEXT,
        resolved_img3 TEXT,
        resolved_img4 TEXT,
        resolved_img5 TEXT,
        resolved_img6 TEXT,
        status TEXT DEFAULT 'Pending Review',
        resolution_notes TEXT,
        imported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        resolved_at TIMESTAMP
    );
    """)
    conn.commit()
    conn.close()

def categorize_error(reason):
    if not reason:
        return "UNKNOWN"
    r = str(reason).lower()
    if "image data issue" in r or "replace the image url" in r:
        return "IMAGE_DATA_ISSUE"
    elif "deleted by the user" in r or "deleted" in r:
        return "PRODUCT_DELETED"
    elif "title" in r:
        return "TITLE_ISSUE"
    elif "category" in r:
        return "CATEGORY_ISSUE"
    elif "price" in r or "moq" in r:
        return "PRICING_ISSUE"
    return "OTHER_ERROR"

def ingest_failed_excel(filepath=None):
    """
    Ingests an Alibaba upload result report Excel file (e.g. 29229346989_377707.xlsx).
    """
    if not filepath:
        filepath = os.path.join(INPUT_DIR, "29229346989_377707.xlsx")
    
    if not os.path.exists(filepath):
        return {"success": False, "error": f"File not found: {filepath}"}

    init_upload_reports_table()
    fname = os.path.basename(filepath)
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb['Alibaba default template'] if 'Alibaba default template' in wb.sheetnames else wb.active

    conn = get_db()
    cur = conn.cursor()

    # Clear prior import of the exact same report file to prevent duplicates
    cur.execute("DELETE FROM alibaba_upload_reports WHERE file_name = ?", (fname,))

    ingested = 0
    for r in range(4, ws.max_row + 1):
        title = ws.cell(r, 1).value
        sku = ws.cell(r, 29).value
        model = ws.cell(r, 28).value
        failed_reason = ws.cell(r, 58).value
        img1 = ws.cell(r, 2).value
        img2 = ws.cell(r, 3).value

        if not sku and not title:
            continue

        sku_str = str(sku).strip() if sku else f"UNKNOWN_ROW_{r}"
        reason_str = str(failed_reason).strip() if failed_reason else "No specific reason provided"
        cat = categorize_error(reason_str)

        cur.execute("""
        INSERT INTO alibaba_upload_reports (
            file_name, file_row, sku, model_number, product_title,
            failed_reason, error_category, original_img1, original_img2,
            status, imported_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'Pending Fix', CURRENT_TIMESTAMP)
        """, (
            fname, r, sku_str, str(model).strip() if model else "",
            str(title).strip() if title else "",
            reason_str, cat,
            str(img1).strip() if img1 else "",
            str(img2).strip() if img2 else ""
        ))
        ingested += 1

    conn.commit()
    conn.close()

    return {"success": True, "file_name": fname, "ingested_count": ingested}

def get_failed_reports_summary():
    """
    Returns statistics and items from the upload reports.
    """
    init_upload_reports_table()
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT COUNT(*) FROM alibaba_upload_reports")
    total = cur.fetchone()[0]

    cur.execute("SELECT error_category, COUNT(*) FROM alibaba_upload_reports GROUP BY error_category")
    by_category = dict(cur.fetchall())

    cur.execute("SELECT status, COUNT(*) FROM alibaba_upload_reports GROUP BY status")
    by_status = dict(cur.fetchall())

    cur.execute("""
    SELECT r.*, p.category_group, p.price_usd
    FROM alibaba_upload_reports r
    LEFT JOIN products p ON r.sku = p.sku
    ORDER BY r.file_row ASC
    """)
    rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    return {
        "total_failed": total,
        "image_data_issues": by_category.get("IMAGE_DATA_ISSUE", 0),
        "product_deleted": by_category.get("PRODUCT_DELETED", 0),
        "other_issues": sum(v for k, v in by_category.items() if k not in ["IMAGE_DATA_ISSUE", "PRODUCT_DELETED"]),
        "status_summary": by_status,
        "items": rows
    }

def export_resolved_excel(target_path=None):
    """
    Exports all resolved items from alibaba_upload_reports into a clean, ready-to-upload Alibaba Excel.
    """
    if not target_path:
        target_path = os.path.join(OUTPUT_DIR, "Alibaba_Bulk_Upload_Resolved_Images.xlsx")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb['Alibaba default template']
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
    SELECT r.*, p.*
    FROM alibaba_upload_reports r
    JOIN products p ON r.sku = p.sku
    WHERE r.status = 'Resolved & Compiled' OR r.resolved_img1 IS NOT NULL
    ORDER BY r.file_row ASC
    """)
    rows = cur.fetchall()

    for idx, r in enumerate(rows, start=1):
        row_idx = 4 + idx
        price = float(r['price_usd'] or 5.50)
        img1 = r['resolved_img1'] or r['image_1'] or ""
        img2 = r['resolved_img2'] or r['image_2'] or ""
        img3 = r['resolved_img3'] or r['image_3'] or ""
        img4 = r['resolved_img4'] or r['image_4'] or ""
        img5 = r['resolved_img5'] or r['image_5'] or ""
        img6 = r['resolved_img6'] or r['image_6'] or ""

        row_data = [
            r['product_title'],
            img1, img2, img3, img4, img5, img6,
            r['description_html'] or f"<p>AP Boots {r['product_title']}. High quality factory export.</p>",
            r['place_of_origin'] or "Indonesia",
            r['brand_name'] or "AP Boots",
            r['category'] or "Shoes & Accessories > Men's Shoes > Boots",
            "Outsole Material", r['outsole_material'] or "PVC / Rubber",
            "Upper Material", r['upper_material'] or "PVC",
            "Season", r['season'] or "All Seasons",
            "Feature", r['feature'] or "100% Waterproof, Anti-Slip, Durable",
            "Gender", r['gender'] or "Men / Unisex",
            r['sell_by'] or "piece",
            None,
            r['unit'] or "Pair/Pairs",
            r['currency'] or "USD",
            r['pricing_type'] or "Tiered pricing",
            r['inventory'] or 5000,
            r['model_number'] or r['sku'],
            r['sku'],
            price,
            "Color", r['color'] or "Standard Export Assorted",
            "Size", r['size_range'] or "EU 38-44",
            None, None, None, None,
            price, 100, round(price*0.95, 2), 500, round(price*0.90, 2), 1000, round(price*0.85, 2), price,
            r['gross_weight_kg'] or 1.7,
            r['shipping_template'] or "Default Shipping Template",
            100, r['lead_time_1'] or 7,
            500, r['lead_time_2'] or 14,
            1000, r['lead_time_3'] or 30,
            r['pkg_length'] or 60.0,
            r['pkg_width'] or 42.0,
            r['pkg_height'] or 38.0
        ]
        for col_idx, val in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    conn.close()
    wb.save(target_path)
    return len(rows), target_path

