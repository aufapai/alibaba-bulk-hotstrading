import http.server
import socketserver
import json
import sqlite3
import os
import re
import urllib.parse
import openpyxl
import sys

if sys.platform.startswith('win'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

BASE_DIR = r"d:\Work\Hots\Product Input"
DB_PATH = os.path.join(BASE_DIR, "Database", "padi_umkm_master.db") if os.path.exists(os.path.join(BASE_DIR, "Database", "padi_umkm_master.db")) else os.path.join(BASE_DIR, "padi_umkm_master.db")
ALIBABA_DB_PATH = os.path.join(BASE_DIR, "Database", "alibaba_master.db") if os.path.exists(os.path.join(BASE_DIR, "Database", "alibaba_master.db")) else os.path.join(BASE_DIR, "alibaba_master.db")
OUTPUT_PADI = os.path.join(BASE_DIR, "Output_PadiUMKM")
PORT = 8889

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def get_alibaba_images():
    """Mengambil katalog foto CDN Tokopedia dari alibaba_master.db untuk Image Picker"""
    if not os.path.exists(ALIBABA_DB_PATH):
        return []
    try:
        conn = sqlite3.connect(ALIBABA_DB_PATH)
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT sku, product_title, image_1, image_2, image_3, image_4, image_5, image_6 FROM products ORDER BY id ASC")
        rows = cur.fetchall()
        conn.close()

        catalog = []
        for r in rows:
            for idx in range(1, 7):
                img_url = r[f'image_{idx}']
                if img_url and str(img_url).startswith('http'):
                    # Shopee: non-main image jangan di-load
                    if 'susercontent.com' in str(img_url) and idx > 1:
                        continue
                    catalog.append({
                        "sku": r['sku'],
                        "title": r['product_title'],
                        "url": img_url,
                        "slot": idx
                    })
        return catalog
    except Exception as e:
        print(f"[WARN] Error fetching alibaba images: {e}")
        return []

def patch_excel_row(batch_num, row_in_batch, prod_data):
    """Menulis update langsung ke file batch Excel PaDi UMKM tertentu"""
    batch_file = os.path.join(OUTPUT_PADI, f"PadiUMKM_Bulk_Upload_Batch_{batch_num}.xlsx")
    if not os.path.exists(batch_file):
        raise FileNotFoundError(f"File batch tidak ditemukan: {batch_file}")

    wb = openpyxl.load_workbook(batch_file)
    ws = wb['Isian Template Upload Barang']
    r = int(row_in_batch)

    # 33 Kolom PaDi UMKM Standard
    ws.cell(r, 1, None)                                        # Col 1: Error_Message
    ws.cell(r, 2, prod_data.get('tipe_produk', 'Barang'))      # Col 2: Tipe_Produk
    ws.cell(r, 3, prod_data.get('nama_produk'))                 # Col 3: Nama_Produk (Maks 70 char)
    ws.cell(r, 4, prod_data.get('deskripsi_produk'))            # Col 4: Deskripsi_Produk (Min 50 char)
    ws.cell(r, 5, int(prod_data.get('kategori_kode', 2450)))   # Col 5: Kategori_Kode
    ws.cell(r, 6, int(prod_data.get('berat_produk_gr', 1000)))  # Col 6: Berat_Produk_gr
    ws.cell(r, 7, int(prod_data.get('dimensi_panjang_cm', 30)))# Col 7: Dimensi_Panjang_cm
    ws.cell(r, 8, int(prod_data.get('dimensi_lebar_cm', 20)))  # Col 8: Dimensi_Lebar_cm
    ws.cell(r, 9, int(prod_data.get('dimensi_tinggi_cm', 15))) # Col 9: Dimensi_Tinggi_cm
    ws.cell(r, 10, int(prod_data.get('pembelian_minimum', 1))) # Col 10: Pembelian_Minimum
    ws.cell(r, 11, prod_data.get('brand', 'AP Boots'))         # Col 11: Brand
    ws.cell(r, 12, prod_data.get('pdn', 'Ya'))                 # Col 12: PDN
    ws.cell(r, 13, prod_data.get('jenis_tkdn', 'Tidak Mengetahui')) # Col 13: Jenis_TKDN
    ws.cell(r, 14, float(prod_data['nilai_tkdn']) if prod_data.get('nilai_tkdn') else None) # Col 14
    ws.cell(r, 15, prod_data.get('dokumen_pendukung_tkdn') or None) # Col 15
    ws.cell(r, 16, prod_data.get('foto_produk_1') or None)     # Col 16: Foto_Produk_1
    ws.cell(r, 17, prod_data.get('foto_produk_2') or None)     # Col 17: Foto_Produk_2
    ws.cell(r, 18, prod_data.get('foto_produk_3') or None)     # Col 18: Foto_Produk_3
    ws.cell(r, 19, prod_data.get('foto_produk_4') or None)     # Col 19: Foto_Produk_4
    ws.cell(r, 20, prod_data.get('foto_produk_5') or None)     # Col 20: Foto_Produk_5
    ws.cell(r, 21, prod_data.get('url_video_produk_1') or None)# Col 21: URL_Video_Produk_1
    ws.cell(r, 22, int(prod_data.get('stok_unit', 100)))       # Col 22: Stok_Unit
    ws.cell(r, 23, prod_data.get('satuan_stok', 'pcs'))        # Col 23: Satuan_Stok
    ws.cell(r, 24, prod_data.get('nama_satuan_lainnya') or None)# Col 24: Nama_Satuan_Lainnya
    ws.cell(r, 25, prod_data.get('stok_selalu_ada', 'Aktif'))  # Col 25: Stok_Selalu_Ada
    ws.cell(r, 26, int(prod_data['pre_order_duration']) if prod_data.get('pre_order_duration') else None) # Col 26
    ws.cell(r, 27, int(prod_data.get('harga_satuan', 0)))      # Col 27: Harga_Satuan
    ws.cell(r, 28, prod_data.get('ppn', 'Non Aktif'))          # Col 28: PPN
    ws.cell(r, 29, prod_data.get('opsi_pph', 'PPH 22'))        # Col 29: Opsi_PPh
    ws.cell(r, 30, prod_data.get('asuransi_pengiriman', 'Wajib')) # Col 30: Asuransi_Pengiriman
    ws.cell(r, 31, prod_data.get('jadikan_gratis_ongkir', 'Tidak')) # Col 31: Gratis Ongkir
    ws.cell(r, 32, prod_data.get('visibilitas_produk', 'Publik')) # Col 32: Visibilitas
    ws.cell(r, 33, prod_data.get('kode_sku'))                  # Col 33: Kode_SKU

    wb.save(batch_file)
    print(f"[EXCEL SYNC] Batch {batch_num} Row {r} updated successfully in '{batch_file}'")

def export_all_batches_from_db():
    """Sinkronisasi seluruh 4 batch Excel langsung dari database SQLite"""
    conn = get_db()
    cur = conn.cursor()
    total_synced = 0

    for b_num in range(1, 5):
        bf = os.path.join(OUTPUT_PADI, f"PadiUMKM_Bulk_Upload_Batch_{b_num}.xlsx")
        if not os.path.exists(bf):
            continue
        wb = openpyxl.load_workbook(bf)
        ws = wb['Isian Template Upload Barang']
        
        cur.execute("SELECT * FROM padi_products WHERE batch_num = ? ORDER BY row_in_batch ASC", (b_num,))
        rows = cur.fetchall()

        for r_data in rows:
            r = r_data['row_in_batch']
            ws.cell(r, 1, None)
            ws.cell(r, 2, r_data['tipe_produk'])
            ws.cell(r, 3, r_data['nama_produk'])
            ws.cell(r, 4, r_data['deskripsi_produk'])
            ws.cell(r, 5, r_data['kategori_kode'])
            ws.cell(r, 6, r_data['berat_produk_gr'])
            ws.cell(r, 7, r_data['dimensi_panjang_cm'])
            ws.cell(r, 8, r_data['dimensi_lebar_cm'])
            ws.cell(r, 9, r_data['dimensi_tinggi_cm'])
            ws.cell(r, 10, r_data['pembelian_minimum'])
            ws.cell(r, 11, r_data['brand'])
            ws.cell(r, 12, r_data['pdn'])
            ws.cell(r, 13, r_data['jenis_tkdn'])
            ws.cell(r, 14, r_data['nilai_tkdn'])
            ws.cell(r, 15, r_data['dokumen_pendukung_tkdn'])
            ws.cell(r, 16, r_data['foto_produk_1'] or None)
            ws.cell(r, 17, r_data['foto_produk_2'] or None)
            ws.cell(r, 18, r_data['foto_produk_3'] or None)
            ws.cell(r, 19, r_data['foto_produk_4'] or None)
            ws.cell(r, 20, r_data['foto_produk_5'] or None)
            ws.cell(r, 21, r_data['url_video_produk_1'] or None)
            ws.cell(r, 22, r_data['stok_unit'])
            ws.cell(r, 23, r_data['satuan_stok'])
            ws.cell(r, 24, r_data['nama_satuan_lainnya'])
            ws.cell(r, 25, r_data['stok_selalu_ada'])
            ws.cell(r, 26, r_data['pre_order_duration'])
            ws.cell(r, 27, r_data['harga_satuan'])
            ws.cell(r, 28, r_data['ppn'])
            ws.cell(r, 29, r_data['opsi_pph'])
            ws.cell(r, 30, r_data['asuransi_pengiriman'])
            ws.cell(r, 31, r_data['jadikan_gratis_ongkir'])
            ws.cell(r, 32, r_data['visibilitas_produk'])
            ws.cell(r, 33, r_data['kode_sku'])
            total_synced += 1

        wb.save(bf)

    conn.close()
    return total_synced

HTML_PAGE = """<!DOCTYPE html>
<html lang="id">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>PaDi UMKM B2B Listing Manager (2026) — PT Indo Sukses Semesta</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
    <style>
        :root {
            --bg-base: #0b1120;
            --bg-card: #131c31;
            --bg-input: #1a243c;
            --border: #243252;
            --border-hover: #3b82f6;
            --text-main: #f8fafc;
            --text-muted: #94a3b8;
            --primary: #0284c7;
            --primary-hover: #0369a1;
            --accent: #10b981;
            --accent-hover: #059669;
            --warning: #f59e0b;
            --danger: #ef4444;
            --padi-teal: #06b6d4;
            --shadow-sm: 0 2px 8px rgba(0,0,0,0.3);
            --shadow-lg: 0 12px 28px rgba(0,0,0,0.45);
        }

        * { box-sizing: border-box; margin: 0; padding: 0; font-family: 'Plus Jakarta Sans', sans-serif; }
        body { background-color: var(--bg-base); color: var(--text-main); min-height: 100vh; display: flex; flex-direction: column; }

        /* Top Header */
        header {
            background: linear-gradient(135deg, #091733 0%, #0c2b5e 100%);
            border-bottom: 1px solid var(--border);
            padding: 1rem 2rem;
            position: sticky;
            top: 0;
            z-index: 50;
            display: flex;
            justify-content: space-between;
            align-items: center;
            box-shadow: var(--shadow-sm);
        }
        .brand-container { display: flex; align-items: center; gap: 14px; }
        .brand-logo {
            width: 44px; height: 44px; border-radius: 10px;
            background: linear-gradient(135deg, #0284c7, #10b981);
            display: flex; align-items: center; justify-content: center;
            font-weight: 800; font-size: 1.25rem; color: #fff;
            box-shadow: 0 4px 12px rgba(2, 132, 199, 0.4);
        }
        .brand-info h1 { font-size: 1.25rem; font-weight: 800; color: #fff; display: flex; align-items: center; gap: 8px; }
        .brand-info p { font-size: 0.8rem; color: var(--text-muted); }
        .tag-badge {
            background: rgba(16, 185, 129, 0.15); color: #34d399;
            border: 1px solid rgba(16, 185, 129, 0.3);
            font-size: 0.7rem; font-weight: 700; padding: 2px 8px; border-radius: 6px;
        }

        .header-actions { display: flex; align-items: center; gap: 10px; }
        .btn {
            display: inline-flex; align-items: center; gap: 8px;
            padding: 9px 16px; border-radius: 8px; font-size: 0.85rem; font-weight: 600;
            cursor: pointer; transition: all 0.2s; border: none; outline: none; text-decoration: none;
        }
        .btn-success { background: linear-gradient(135deg, #10b981, #059669); color: #fff; }
        .btn-success:hover { filter: brightness(1.1); transform: translateY(-1px); }
        .btn-primary { background: linear-gradient(135deg, #0284c7, #0369a1); color: #fff; }
        .btn-primary:hover { filter: brightness(1.1); }
        .btn-outline { background: transparent; border: 1px solid var(--border); color: var(--text-muted); }
        .btn-outline:hover { background: rgba(255,255,255,0.05); color: #fff; border-color: var(--text-muted); }

        /* Container */
        .main-container { max-width: 1480px; width: 100%; margin: 0 auto; padding: 1.5rem 2rem; flex: 1; }

        /* Stats Section */
        .stats-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
            gap: 1rem;
            margin-bottom: 1.5rem;
        }
        .stat-card {
            background: var(--bg-card);
            border: 1px solid var(--border);
            border-radius: 12px;
            padding: 1.15rem;
            position: relative;
            overflow: hidden;
        }
        .stat-card::after {
            content: ''; position: absolute; top: 0; left: 0; right: 0; height: 3px;
            background: linear-gradient(90deg, var(--primary), var(--accent));
        }
        .stat-title { font-size: 0.75rem; text-transform: uppercase; letter-spacing: 0.05em; color: var(--text-muted); font-weight: 600; }
        .stat-value { font-size: 1.7rem; font-weight: 800; color: #fff; margin-top: 4px; }
        .stat-desc { font-size: 0.75rem; color: #34d399; margin-top: 4px; font-weight: 500; }

        /* Batch Navigation & Toolbar */
        .control-panel {
            background: var(--bg-card);
            border: 1px solid var(--border);
            border-radius: 12px;
            padding: 1rem 1.25rem;
            margin-bottom: 1.5rem;
            display: flex;
            flex-direction: column;
            gap: 1rem;
        }
        .batch-tabs { display: flex; gap: 8px; flex-wrap: wrap; }
        .batch-tab {
            padding: 8px 16px; border-radius: 8px; font-size: 0.825rem; font-weight: 600;
            background: var(--bg-input); color: var(--text-muted); border: 1px solid var(--border);
            cursor: pointer; transition: all 0.2s; display: flex; align-items: center; gap: 6px;
        }
        .batch-tab:hover { background: #223150; color: #fff; }
        .batch-tab.active {
            background: var(--primary); color: #fff; border-color: var(--primary-hover);
            box-shadow: 0 4px 12px rgba(2, 132, 199, 0.35);
        }
        .batch-badge {
            background: rgba(0,0,0,0.25); padding: 2px 7px; border-radius: 999px; font-size: 0.7rem;
        }

        .search-row { display: flex; gap: 1rem; flex-wrap: wrap; align-items: center; justify-content: space-between; }
        .search-box { flex: 1; min-width: 300px; position: relative; }
        .search-box input {
            width: 100%; background: var(--bg-input); border: 1px solid var(--border);
            border-radius: 8px; padding: 10px 14px 10px 38px; color: #fff; font-size: 0.875rem; outline: none;
        }
        .search-box input:focus { border-color: var(--primary); }
        .search-icon { position: absolute; left: 12px; top: 50%; transform: translateY(-50%); font-size: 0.9rem; color: var(--text-muted); }

        .batch-download-group { display: flex; gap: 6px; align-items: center; flex-wrap: wrap; }
        .download-btn {
            font-size: 0.75rem; padding: 6px 12px; border-radius: 6px; background: rgba(59, 130, 246, 0.12);
            color: #60a5fa; border: 1px solid rgba(59, 130, 246, 0.25); text-decoration: none; font-weight: 600;
            transition: all 0.2s;
        }
        .download-btn:hover { background: rgba(59, 130, 246, 0.25); color: #fff; }

        /* Products Grid */
        .products-grid {
            display: grid;
            grid-template-columns: repeat(auto-fill, minmax(340px, 1fr));
            gap: 1.25rem;
        }
        .product-card {
            background: var(--bg-card);
            border: 1px solid var(--border);
            border-radius: 12px;
            overflow: hidden;
            display: flex;
            flex-direction: column;
            transition: transform 0.2s, border-color 0.2s;
        }
        .product-card:hover { transform: translateY(-3px); border-color: var(--border-hover); }

        .card-top { position: relative; height: 190px; background: #080d19; display: flex; align-items: center; justify-content: center; overflow: hidden; }
        .card-top img { width: 100%; height: 100%; object-fit: contain; padding: 12px; transition: transform 0.3s; }
        .product-card:hover .card-top img { transform: scale(1.05); }
        .no-img { color: var(--text-muted); font-size: 0.85rem; display: flex; flex-direction: column; align-items: center; gap: 4px; }

        .card-badges { position: absolute; top: 10px; left: 10px; display: flex; gap: 6px; }
        .badge-batch { background: rgba(2, 132, 199, 0.85); color: #fff; font-size: 0.65rem; font-weight: 700; padding: 3px 8px; border-radius: 4px; backdrop-filter: blur(4px); }
        .badge-sku { background: rgba(19, 28, 49, 0.85); border: 1px solid var(--border); color: #38bdf8; font-family: 'JetBrains Mono', monospace; font-size: 0.7rem; font-weight: 600; padding: 3px 8px; border-radius: 4px; backdrop-filter: blur(4px); }

        .card-gallery-count {
            position: absolute; bottom: 8px; right: 10px;
            background: rgba(0,0,0,0.65); color: #e2e8f0; font-size: 0.65rem;
            padding: 2px 7px; border-radius: 4px; backdrop-filter: blur(4px);
        }

        .card-body { padding: 1.15rem; display: flex; flex-direction: column; flex: 1; }
        .card-title {
            font-size: 0.925rem; font-weight: 700; color: #fff; line-height: 1.35;
            margin-bottom: 8px; min-height: 2.7em; display: -webkit-box; -webkit-line-clamp: 2; -webkit-box-orient: vertical; overflow: hidden;
        }
        
        .char-counter {
            font-size: 0.7rem; font-weight: 600; display: inline-flex; align-items: center; gap: 4px; margin-bottom: 10px;
        }
        .char-ok { color: #34d399; }
        .char-warn { color: #f87171; }

        .card-price-row {
            display: flex; justify-content: space-between; align-items: baseline;
            margin-top: auto; padding-top: 10px; border-top: 1px solid var(--border);
        }
        .price-label { font-size: 0.725rem; color: var(--text-muted); }
        .price-val { font-size: 1.15rem; font-weight: 800; color: #38bdf8; font-family: 'JetBrains Mono', monospace; }

        .card-footer { padding: 0.85rem 1.15rem; background: rgba(0,0,0,0.15); border-top: 1px solid var(--border); display: flex; gap: 8px; }
        .btn-card { flex: 1; justify-content: center; padding: 8px; font-size: 0.8rem; }

        /* Modal Edit */
        .modal-overlay {
            position: fixed; inset: 0; background: rgba(0,0,0,0.8); backdrop-filter: blur(6px);
            z-index: 100; display: none; align-items: center; justify-content: center; padding: 1.5rem;
        }
        .modal-overlay.active { display: flex; }
        .modal-box {
            background: var(--bg-card); border: 1px solid var(--border); border-radius: 16px;
            width: 100%; max-width: 980px; max-height: 90vh; display: flex; flex-direction: column;
            box-shadow: var(--shadow-lg); overflow: hidden;
        }
        .modal-header {
            padding: 1.25rem 1.75rem; border-bottom: 1px solid var(--border);
            display: flex; justify-content: space-between; align-items: center;
            background: linear-gradient(135deg, #0f172a, #1e293b);
        }
        .modal-header h3 { font-size: 1.15rem; font-weight: 800; color: #fff; }
        .btn-close { background: transparent; border: none; color: var(--text-muted); font-size: 1.4rem; cursor: pointer; }
        .btn-close:hover { color: #fff; }

        .modal-body { padding: 1.5rem 1.75rem; overflow-y: auto; flex: 1; display: flex; flex-direction: column; gap: 1.5rem; }

        /* Section Title */
        .form-section-title {
            font-size: 0.85rem; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em;
            color: var(--padi-teal); margin-bottom: 0.85rem; display: flex; align-items: center; gap: 8px;
        }
        .form-section-title::after { content: ''; flex: 1; height: 1px; background: var(--border); }

        /* Image Slots Manager */
        .image-slots-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
            gap: 1rem;
        }
        .img-slot-card {
            background: var(--bg-input); border: 1px solid var(--border); border-radius: 10px;
            padding: 10px; display: flex; flex-direction: column; gap: 8px;
        }
        .img-slot-header { display: flex; justify-content: space-between; align-items: center; font-size: 0.725rem; font-weight: 700; color: var(--text-muted); }
        .slot-thumb {
            width: 100%; height: 120px; background: #080d19; border-radius: 6px;
            display: flex; align-items: center; justify-content: center; overflow: hidden; border: 1px solid var(--border);
        }
        .slot-thumb img { width: 100%; height: 100%; object-fit: contain; }
        .slot-input {
            width: 100%; background: #0b1120; border: 1px solid var(--border);
            border-radius: 6px; padding: 6px 8px; color: #fff; font-size: 0.725rem; outline: none;
            font-family: 'JetBrains Mono', monospace;
        }
        .slot-actions { display: flex; gap: 4px; }
        .btn-slot {
            flex: 1; padding: 5px; font-size: 0.7rem; border-radius: 4px;
            background: #243252; color: #e2e8f0; border: none; cursor: pointer; text-align: center;
        }
        .btn-slot:hover { background: #3b82f6; color: #fff; }
        .btn-slot-clear { background: rgba(239, 68, 68, 0.15); color: #f87171; border: 1px solid rgba(239, 68, 68, 0.25); }
        .btn-slot-clear:hover { background: #ef4444; color: #fff; }

        /* Form Inputs */
        .form-row { display: grid; grid-template-columns: repeat(auto-fit, minmax(240px, 1fr)); gap: 1rem; }
        .form-group { display: flex; flex-direction: column; gap: 6px; }
        .form-label { font-size: 0.775rem; font-weight: 600; color: var(--text-muted); display: flex; justify-content: space-between; }
        .form-input, .form-textarea, .form-select {
            width: 100%; background: var(--bg-input); border: 1px solid var(--border);
            border-radius: 8px; padding: 9px 12px; color: #fff; font-size: 0.85rem; outline: none;
        }
        .form-input:focus, .form-textarea:focus, .form-select:focus { border-color: var(--primary); }
        .form-textarea { resize: vertical; min-height: 90px; }

        .modal-footer {
            padding: 1.15rem 1.75rem; border-top: 1px solid var(--border);
            background: #0f172a; display: flex; justify-content: space-between; align-items: center;
        }

        /* Image Picker Drawer */
        .picker-modal {
            position: fixed; inset: 0; background: rgba(0,0,0,0.85); backdrop-filter: blur(8px);
            z-index: 150; display: none; align-items: center; justify-content: center; padding: 1.5rem;
        }
        .picker-modal.active { display: flex; }
        .picker-box {
            background: var(--bg-card); border: 1px solid var(--border); border-radius: 16px;
            width: 100%; max-width: 900px; max-height: 85vh; display: flex; flex-direction: column; overflow: hidden;
        }
        .picker-header { padding: 1rem 1.5rem; border-bottom: 1px solid var(--border); display: flex; justify-content: space-between; align-items: center; }
        .picker-search { padding: 0.75rem 1.5rem; border-bottom: 1px solid var(--border); background: #0b1120; }
        .picker-grid { padding: 1.25rem; overflow-y: auto; display: grid; grid-template-columns: repeat(auto-fill, minmax(130px, 1fr)); gap: 10px; flex: 1; }
        .picker-item {
            background: var(--bg-input); border: 1px solid var(--border); border-radius: 8px;
            padding: 8px; cursor: pointer; display: flex; flex-direction: column; align-items: center; gap: 6px;
            transition: all 0.2s;
        }
        .picker-item:hover { border-color: var(--primary); transform: scale(1.03); background: #223150; }
        .picker-item img { width: 100px; height: 100px; object-fit: contain; }
        .picker-item span { font-size: 0.65rem; color: var(--text-muted); text-align: center; line-height: 1.2; word-break: break-all; }

        /* Toast */
        .toast {
            position: fixed; bottom: 24px; right: 24px; z-index: 200;
            background: #10b981; color: #fff; padding: 12px 20px; border-radius: 8px;
            font-size: 0.85rem; font-weight: 600; box-shadow: var(--shadow-lg);
            display: none; align-items: center; gap: 10px; animation: slideUp 0.3s ease;
        }
        @keyframes slideUp { from { transform: translateY(20px); opacity: 0; } to { transform: translateY(0); opacity: 1; } }
    </style>
</head>
<body>

    <!-- Header -->
    <header>
        <div class="brand-container">
            <div class="brand-logo">P</div>
            <div class="brand-info">
                <h1>PaDi UMKM Manager <span class="tag-badge">B2B BUMN 2026</span></h1>
                <p>PT Indo Sukses Semesta — Official AP Boots Catalog & Bulk Upload Engine</p>
            </div>
        </div>
        <div class="header-actions">
            <button class="btn btn-success" onclick="exportAllBatches()" id="btn-sync-all">
                <span>⚡ Sinkronisasi & Export Semua Excel</span>
            </button>
            <a href="http://localhost:8888" target="_blank" class="btn btn-outline" title="Buka Alibaba Manager">
                <span>🌐 Alibaba Manager</span>
            </a>
        </div>
    </header>

    <!-- Main Container -->
    <div class="main-container">

        <!-- Stats Overview -->
        <div class="stats-grid">
            <div class="stat-card">
                <div class="stat-title">Total Produk Siap Upload</div>
                <div class="stat-value" id="stat-total">76</div>
                <div class="stat-desc">4 Batch File Excel Terverifikasi</div>
            </div>
            <div class="stat-card">
                <div class="stat-title">Total Batch Pengadaan</div>
                <div class="stat-value">4 Batch</div>
                <div class="stat-desc">Maksimal 20 SKU per File Sesuai Regulasi</div>
            </div>
            <div class="stat-card">
                <div class="stat-title">Status Gambar CDN</div>
                <div class="stat-value" id="stat-img-status">100% Aktif</div>
                <div class="stat-desc">Tokopedia High-Res CDN Tanpa Duplikasi</div>
            </div>
            <div class="stat-card">
                <div class="stat-title">Kepatuhan Kolom</div>
                <div class="stat-value" style="color: #34d399;">33 Kolom Standard</div>
                <div class="stat-desc">PDN 100% & PPh Pasal 22 Terkonfigurasi</div>
            </div>
        </div>

        <!-- Controls & Batch Navigation -->
        <div class="control-panel">
            <div class="batch-tabs">
                <button class="batch-tab active" onclick="setBatchFilter(0)">
                    <span>Semua Batch</span> <span class="batch-badge" id="badge-all">76</span>
                </button>
                <button class="batch-tab" onclick="setBatchFilter(1)">
                    <span>Batch 1 (Safety & Casual)</span> <span class="batch-badge">20</span>
                </button>
                <button class="batch-tab" onclick="setBatchFilter(2)">
                    <span>Batch 2 (Series 9 & Eco)</span> <span class="batch-badge">20</span>
                </button>
                <button class="batch-tab" onclick="setBatchFilter(3)">
                    <span>Batch 3 (Kids, Sandal & Max)</span> <span class="batch-badge">20</span>
                </button>
                <button class="batch-tab" onclick="setBatchFilter(4)">
                    <span>Batch 4 (Roots & Allbike)</span> <span class="batch-badge">16</span>
                </button>
            </div>

            <div class="search-row">
                <div class="search-box">
                    <span class="search-icon">🔍</span>
                    <input type="text" id="searchInput" placeholder="Cari berdasarkan Nama Sepatu, Kode SKU, atau Kategori..." oninput="handleSearch()">
                </div>
                <div class="batch-download-group">
                    <span style="font-size: 0.75rem; color: var(--text-muted); font-weight: 600;">Download File Excel Siap Upload:</span>
                    <a href="/api/download/1" class="download-btn" download>📥 Batch 1 (.xlsx)</a>
                    <a href="/api/download/2" class="download-btn" download>📥 Batch 2 (.xlsx)</a>
                    <a href="/api/download/3" class="download-btn" download>📥 Batch 3 (.xlsx)</a>
                    <a href="/api/download/4" class="download-btn" download>📥 Batch 4 (.xlsx)</a>
                </div>
            </div>
        </div>

        <!-- Products Grid -->
        <div class="products-grid" id="productsGrid">
            <!-- Product Cards loaded via JS -->
        </div>

    </div>

    <!-- Edit Modal -->
    <div class="modal-overlay" id="editModal">
        <div class="modal-box">
            <div class="modal-header">
                <div>
                    <h3 id="modalProductTitle">Edit Produk & Link Gambar</h3>
                    <p style="font-size: 0.75rem; color: var(--text-muted); margin-top: 2px;" id="modalProductSubtitle">Batch X • Baris Y</p>
                </div>
                <button class="btn-close" onclick="closeModal()">&times;</button>
            </div>

            <div class="modal-body">
                <!-- Section 1: Image URLs (Slot 1 - 5) -->
                <div>
                    <div class="form-section-title">🖼️ Pengaturan 5 Foto Produk (Direct CDN URL)</div>
                    <p style="font-size: 0.75rem; color: var(--text-muted); margin-bottom: 12px;">
                        PaDi UMKM memerlukan direct image URL yang aktif dan bisa diakses publik. Gunakan tombol <strong>"Pilih dari CDN"</strong> untuk mengambil foto resmi berlatar putih bersih dari katalog Tokopedia.
                    </p>
                    <div class="image-slots-grid">
                        <!-- Slot 1 -->
                        <div class="img-slot-card">
                            <div class="img-slot-header">
                                <span>Foto 1 (Utama) *</span>
                            </div>
                            <div class="slot-thumb" id="slot-thumb-1">
                                <span class="no-img">Belum Ada</span>
                            </div>
                            <input type="text" class="slot-input" id="input-foto-1" placeholder="URL Foto 1..." oninput="updateSlotPreview(1)">
                            <div class="slot-actions">
                                <button class="btn-slot" onclick="openPicker(1)">📂 Pilih CDN</button>
                                <button class="btn-slot btn-slot-clear" onclick="clearSlot(1)">✕</button>
                            </div>
                        </div>

                        <!-- Slot 2 -->
                        <div class="img-slot-card">
                            <div class="img-slot-header">
                                <span>Foto 2 (Sudut 2)</span>
                            </div>
                            <div class="slot-thumb" id="slot-thumb-2">
                                <span class="no-img">Belum Ada</span>
                            </div>
                            <input type="text" class="slot-input" id="input-foto-2" placeholder="URL Foto 2..." oninput="updateSlotPreview(2)">
                            <div class="slot-actions">
                                <button class="btn-slot" onclick="openPicker(2)">📂 Pilih CDN</button>
                                <button class="btn-slot btn-slot-clear" onclick="clearSlot(2)">✕</button>
                            </div>
                        </div>

                        <!-- Slot 3 -->
                        <div class="img-slot-card">
                            <div class="img-slot-header">
                                <span>Foto 3 (Sol / Outsole)</span>
                            </div>
                            <div class="slot-thumb" id="slot-thumb-3">
                                <span class="no-img">Belum Ada</span>
                            </div>
                            <input type="text" class="slot-input" id="input-foto-3" placeholder="URL Foto 3..." oninput="updateSlotPreview(3)">
                            <div class="slot-actions">
                                <button class="btn-slot" onclick="openPicker(3)">📂 Pilih CDN</button>
                                <button class="btn-slot btn-slot-clear" onclick="clearSlot(3)">✕</button>
                            </div>
                        </div>

                        <!-- Slot 4 -->
                        <div class="img-slot-card">
                            <div class="img-slot-header">
                                <span>Foto 4 (Detail Fleksibel)</span>
                            </div>
                            <div class="slot-thumb" id="slot-thumb-4">
                                <span class="no-img">Belum Ada</span>
                            </div>
                            <input type="text" class="slot-input" id="input-foto-4" placeholder="URL Foto 4..." oninput="updateSlotPreview(4)">
                            <div class="slot-actions">
                                <button class="btn-slot" onclick="openPicker(4)">📂 Pilih CDN</button>
                                <button class="btn-slot btn-slot-clear" onclick="clearSlot(4)">✕</button>
                            </div>
                        </div>

                        <!-- Slot 5 -->
                        <div class="img-slot-card">
                            <div class="img-slot-header">
                                <span>Foto 5 (Box / Kemasan)</span>
                            </div>
                            <div class="slot-thumb" id="slot-thumb-5">
                                <span class="no-img">Belum Ada</span>
                            </div>
                            <input type="text" class="slot-input" id="input-foto-5" placeholder="URL Foto 5..." oninput="updateSlotPreview(5)">
                            <div class="slot-actions">
                                <button class="btn-slot" onclick="openPicker(5)">📂 Pilih CDN</button>
                                <button class="btn-slot btn-slot-clear" onclick="clearSlot(5)">✕</button>
                            </div>
                        </div>
                    </div>
                </div>

                <!-- Section 2: Core Details (Excel Columns) -->
                <div>
                    <div class="form-section-title">📋 Spesifikasi Utama Produk (33 Kolom PaDi UMKM)</div>
                    <div class="form-row">
                        <div class="form-group" style="grid-column: span 2;">
                            <div class="form-label">
                                <span>Nama Produk (Maks 70 Karakter) *</span>
                                <span id="titleCounter" class="char-ok">0 / 70</span>
                            </div>
                            <input type="text" class="form-input" id="editNamaProduk" oninput="checkTitleLength()">
                        </div>
                        <div class="form-group">
                            <label class="form-label">Kode SKU Internal</label>
                            <input type="text" class="form-input" id="editKodeSKU" style="font-family: 'JetBrains Mono', monospace;">
                        </div>
                    </div>

                    <div class="form-row" style="margin-top: 1rem;">
                        <div class="form-group">
                            <label class="form-label">Harga Satuan (IDR Sebelum Pajak) *</label>
                            <input type="number" class="form-input" id="editHargaSatuan" style="font-family: 'JetBrains Mono', monospace;">
                        </div>
                        <div class="form-group">
                            <label class="form-label">Berat Produk (Gram) *</label>
                            <input type="number" class="form-input" id="editBerat">
                        </div>
                        <div class="form-group">
                            <label class="form-label">Stok Siap Kirim (Unit)</label>
                            <input type="number" class="form-input" id="editStok">
                        </div>
                        <div class="form-group">
                            <label class="form-label">Kategori Kode PaDi</label>
                            <input type="number" class="form-input" id="editKategoriKode">
                        </div>
                    </div>

                    <div class="form-row" style="margin-top: 1rem;">
                        <div class="form-group">
                            <label class="form-label">Dimensi P x L x T (cm)</label>
                            <div style="display: flex; gap: 6px;">
                                <input type="number" class="form-input" id="editDimP" placeholder="P">
                                <input type="number" class="form-input" id="editDimL" placeholder="L">
                                <input type="number" class="form-input" id="editDimT" placeholder="T">
                            </div>
                        </div>
                        <div class="form-group">
                            <label class="form-label">Status PDN (Produk Dalam Negeri)</label>
                            <select class="form-select" id="editPDN">
                                <option value="Ya">Ya (100% Buatan Indonesia)</option>
                                <option value="Tidak">Tidak</option>
                            </select>
                        </div>
                        <div class="form-group">
                            <label class="form-label">Jenis TKDN</label>
                            <select class="form-select" id="editJenisTKDN">
                                <option value="Tidak Mengetahui">Tidak Mengetahui</option>
                                <option value="Sertifikasi Pihak Ketiga">Sertifikasi Pihak Ketiga (Kemenperin)</option>
                            </select>
                        </div>
                    </div>

                    <div class="form-group" style="margin-top: 1rem;">
                        <div class="form-label">
                            <span>Deskripsi Produk (Minimal 50 Karakter) *</span>
                            <span id="descCounter" class="char-ok">0 karakter</span>
                        </div>
                        <textarea class="form-textarea" id="editDeskripsi" rows="5" oninput="checkDescLength()"></textarea>
                    </div>
                </div>

            </div>

            <div class="modal-footer">
                <button class="btn btn-outline" onclick="closeModal()">Batal</button>
                <button class="btn btn-success" onclick="saveProductChanges()">
                    <span>💾 Simpan ke Database & Patch Excel Langsung</span>
                </button>
            </div>
        </div>
    </div>

    <!-- Image Picker Modal -->
    <div class="picker-modal" id="pickerModal">
        <div class="picker-box">
            <div class="picker-header">
                <div>
                    <h4 style="color: #fff; font-size: 1.1rem; font-weight: 700;">Galeri CDN Foto Tokopedia (Resmi)</h4>
                    <p style="font-size: 0.75rem; color: var(--text-muted);">Pilih salah satu gambar untuk disalin langsung ke slot foto produk</p>
                </div>
                <button class="btn-close" onclick="closePicker()">&times;</button>
            </div>
            <div class="picker-search">
                <input type="text" id="pickerSearchInput" class="form-input" placeholder="Ketik kata kunci untuk menyaring (misal: terra, chef, s5, allbike, 2003)..." oninput="filterPickerGallery()">
            </div>
            <div class="picker-grid" id="pickerGalleryGrid">
                <!-- Gallery items injected via JS -->
            </div>
        </div>
    </div>

    <!-- Toast Notification -->
    <div class="toast" id="toastMessage">
        <span id="toastIcon">✓</span>
        <span id="toastText">Perubahan berhasil disimpan!</span>
    </div>

    <script>
        let allProducts = [];
        let allCatalogImages = [];
        let currentBatch = 0;
        let activeEditingProductId = null;
        let activeTargetSlot = 1;

        async function init() {
            await loadProducts();
            await loadCatalogImages();
        }

        async function loadProducts() {
            try {
                const res = await fetch('/api/products');
                allProducts = await res.json();
                document.getElementById('stat-total').innerText = allProducts.length;
                document.getElementById('badge-all').innerText = allProducts.length;
                renderProducts();
            } catch (err) {
                console.error("Gagal memuat produk:", err);
            }
        }

        async function loadCatalogImages() {
            try {
                const res = await fetch('/api/images/catalog');
                allCatalogImages = await res.json();
            } catch (err) {
                console.error("Gagal memuat galeri:", err);
            }
        }

        function setBatchFilter(batchNum) {
            currentBatch = batchNum;
            const tabs = document.querySelectorAll('.batch-tab');
            tabs.forEach((tab, idx) => {
                if (idx === batchNum) tab.classList.add('active');
                else tab.classList.remove('active');
            });
            renderProducts();
        }

        function handleSearch() {
            renderProducts();
        }

        function renderProducts() {
            const grid = document.getElementById('productsGrid');
            const search = document.getElementById('searchInput').value.toLowerCase().trim();

            let filtered = allProducts.filter(p => {
                const matchBatch = (currentBatch === 0) || (p.batch_num === currentBatch);
                const matchSearch = (!search) || 
                    (p.nama_produk && p.nama_produk.toLowerCase().includes(search)) ||
                    (p.kode_sku && p.kode_sku.toLowerCase().includes(search));
                return matchBatch && matchSearch;
            });

            if (filtered.length === 0) {
                grid.innerHTML = `<div style="grid-column: 1/-1; text-align: center; padding: 3rem; color: var(--text-muted);">Tidak ada produk yang cocok dengan pencarian.</div>`;
                return;
            }

            let html = '';
            filtered.forEach(p => {
                const mainImg = p.foto_produk_1;
                const titleLen = (p.nama_produk || '').length;
                const isTitleOk = titleLen <= 70;
                
                // Count active photos
                let photoCount = 0;
                if (p.foto_produk_1) photoCount++;
                if (p.foto_produk_2) photoCount++;
                if (p.foto_produk_3) photoCount++;
                if (p.foto_produk_4) photoCount++;
                if (p.foto_produk_5) photoCount++;

                const formattedPrice = new Intl.NumberFormat('id-ID', { style: 'currency', currency: 'IDR', maximumFractionDigits: 0 }).format(p.harga_satuan || 0);

                html += `
                <div class="product-card">
                    <div class="card-top">
                        ${mainImg ? `<img src="${mainImg}" alt="${p.nama_produk}" loading="lazy" onerror="this.onerror=null; this.src='https://placehold.co/400x400/1e293b/94a3b8?text=Image+404';">` : `<div class="no-img"><span>📷</span><span>Tidak ada foto</span></div>`}
                        <div class="card-badges">
                            <span class="badge-batch">Batch ${p.batch_num}</span>
                            <span class="badge-sku">${p.kode_sku || 'NO-SKU'}</span>
                        </div>
                        <div class="card-gallery-count">📸 ${photoCount}/5 Foto</div>
                    </div>
                    <div class="card-body">
                        <div class="card-title" title="${p.nama_produk}">${p.nama_produk}</div>
                        <div class="char-counter ${isTitleOk ? 'char-ok' : 'char-warn'}">
                            ${isTitleOk ? '✓' : '⚠️'} ${titleLen} / 70 Karakter PaDi
                        </div>
                        <div class="card-price-row">
                            <span class="price-label">Harga Satuan B2B</span>
                            <span class="price-val">${formattedPrice}</span>
                        </div>
                    </div>
                    <div class="card-footer">
                        <button class="btn btn-primary btn-card" onclick="openEditModal(${p.id})">
                            <span>✏️ Edit & Ganti Foto</span>
                        </button>
                    </div>
                </div>
                `;
            });

            grid.innerHTML = html;
        }

        async function openEditModal(productId) {
            activeEditingProductId = productId;
            const res = await fetch(`/api/products/${productId}`);
            const p = await res.json();

            document.getElementById('modalProductTitle').innerText = p.nama_produk;
            document.getElementById('modalProductSubtitle').innerText = `Batch ${p.batch_num} • Baris Excel ${p.row_in_batch} • SKU: ${p.kode_sku}`;

            // Populate Foto Slots
            for (let i = 1; i <= 5; i++) {
                const url = p[`foto_produk_${i}`] || '';
                document.getElementById(`input-foto-${i}`).value = url;
                updateSlotPreview(i);
            }

            // Populate Core Fields
            document.getElementById('editNamaProduk').value = p.nama_produk || '';
            document.getElementById('editKodeSKU').value = p.kode_sku || '';
            document.getElementById('editHargaSatuan').value = p.harga_satuan || 0;
            document.getElementById('editBerat').value = p.berat_produk_gr || 1000;
            document.getElementById('editStok').value = p.stok_unit || 100;
            document.getElementById('editKategoriKode').value = p.kategori_kode || 2450;
            document.getElementById('editDimP').value = p.dimensi_panjang_cm || 30;
            document.getElementById('editDimL').value = p.dimensi_lebar_cm || 20;
            document.getElementById('editDimT').value = p.dimensi_tinggi_cm || 15;
            document.getElementById('editPDN').value = p.pdn || 'Ya';
            document.getElementById('editJenisTKDN').value = p.jenis_tkdn || 'Tidak Mengetahui';
            document.getElementById('editDeskripsi').value = p.deskripsi_produk || '';

            checkTitleLength();
            checkDescLength();

            document.getElementById('editModal').classList.add('active');
        }

        function closeModal() {
            document.getElementById('editModal').classList.remove('active');
        }

        function updateSlotPreview(slot) {
            const url = document.getElementById(`input-foto-${slot}`).value.trim();
            const thumb = document.getElementById(`slot-thumb-${slot}`);
            if (url) {
                thumb.innerHTML = `<img src="${url}" onerror="this.onerror=null; this.src='https://placehold.co/400x400/1e293b/94a3b8?text=Image+404';">`;
            } else {
                thumb.innerHTML = `<span class="no-img">Belum Ada</span>`;
            }
        }

        function clearSlot(slot) {
            document.getElementById(`input-foto-${slot}`).value = '';
            updateSlotPreview(slot);
        }

        function checkTitleLength() {
            const val = document.getElementById('editNamaProduk').value;
            const counter = document.getElementById('titleCounter');
            counter.innerText = `${val.length} / 70`;
            if (val.length > 70) {
                counter.className = 'char-counter char-warn';
            } else {
                counter.className = 'char-counter char-ok';
            }
        }

        function checkDescLength() {
            const val = document.getElementById('editDeskripsi').value;
            const counter = document.getElementById('descCounter');
            counter.innerText = `${val.length} karakter (Min 50)`;
            if (val.length < 50) {
                counter.className = 'char-counter char-warn';
            } else {
                counter.className = 'char-counter char-ok';
            }
        }

        async function saveProductChanges() {
            if (!activeEditingProductId) return;

            const payload = {
                nama_produk: document.getElementById('editNamaProduk').value.trim(),
                kode_sku: document.getElementById('editKodeSKU').value.trim(),
                harga_satuan: parseInt(document.getElementById('editHargaSatuan').value) || 0,
                berat_produk_gr: parseInt(document.getElementById('editBerat').value) || 1000,
                stok_unit: parseInt(document.getElementById('editStok').value) || 100,
                kategori_kode: parseInt(document.getElementById('editKategoriKode').value) || 2450,
                dimensi_panjang_cm: parseInt(document.getElementById('editDimP').value) || 30,
                dimensi_lebar_cm: parseInt(document.getElementById('editDimL').value) || 20,
                dimensi_tinggi_cm: parseInt(document.getElementById('editDimT').value) || 15,
                pdn: document.getElementById('editPDN').value,
                jenis_tkdn: document.getElementById('editJenisTKDN').value,
                deskripsi_produk: document.getElementById('editDeskripsi').value.trim(),
                foto_produk_1: document.getElementById('input-foto-1').value.trim(),
                foto_produk_2: document.getElementById('input-foto-2').value.trim(),
                foto_produk_3: document.getElementById('input-foto-3').value.trim(),
                foto_produk_4: document.getElementById('input-foto-4').value.trim(),
                foto_produk_5: document.getElementById('input-foto-5').value.trim()
            };

            try {
                const res = await fetch(`/api/products/${activeEditingProductId}`, {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify(payload)
                });
                const result = await res.json();
                if (result.success) {
                    showToast("✓ Sukses! Data tersimpan di SQLite & File Excel Batch ter-patch langsung!");
                    closeModal();
                    await loadProducts();
                } else {
                    alert("Gagal menyimpan: " + (result.error || 'Unknown error'));
                }
            } catch (err) {
                alert("Error koneksi saat menyimpan: " + err);
            }
        }

        async function exportAllBatches() {
            const btn = document.getElementById('btn-sync-all');
            const originalText = btn.innerHTML;
            btn.innerHTML = `<span>⏳ Sedang menyinkronkan 4 file Excel...</span>`;
            btn.style.opacity = '0.7';

            try {
                const res = await fetch('/api/export', { method: 'POST' });
                const result = await res.json();
                if (result.success) {
                    showToast(`✓ Sukses sinkronisasi! ${result.count} produk di 4 Batch Excel ter-update rapi!`);
                } else {
                    alert("Export gagal: " + result.error);
                }
            } catch (err) {
                alert("Error saat export: " + err);
            } finally {
                btn.innerHTML = originalText;
                btn.style.opacity = '1';
            }
        }

        /* Image Picker Logic */
        function openPicker(slot) {
            activeTargetSlot = slot;
            renderPickerGallery();
            document.getElementById('pickerModal').classList.add('active');
        }

        function closePicker() {
            document.getElementById('pickerModal').classList.remove('active');
        }

        function renderPickerGallery() {
            const grid = document.getElementById('pickerGalleryGrid');
            const search = document.getElementById('pickerSearchInput').value.toLowerCase().trim();

            let filtered = allCatalogImages.filter(img => {
                if (!search) return true;
                return (img.title && img.title.toLowerCase().includes(search)) ||
                       (img.sku && img.sku.toLowerCase().includes(search));
            });

            if (filtered.length === 0) {
                grid.innerHTML = `<div style="grid-column: 1/-1; text-align: center; padding: 2rem; color: var(--text-muted);">Tidak ada gambar yang cocok.</div>`;
                return;
            }

            let html = '';
            filtered.forEach(img => {
                html += `
                <div class="picker-item" onclick="selectPickerImage('${img.url}')">
                    <img src="${img.url}" loading="lazy" onerror="this.onerror=null; this.src='https://placehold.co/100x100/1e293b/94a3b8?text=Image';">
                    <span>${img.sku} (${img.slot})</span>
                </div>
                `;
            });
            grid.innerHTML = html;
        }

        function filterPickerGallery() {
            renderPickerGallery();
        }

        function selectPickerImage(url) {
            document.getElementById(`input-foto-${activeTargetSlot}`).value = url;
            updateSlotPreview(activeTargetSlot);
            closePicker();
            showToast(`✓ Foto disalin ke Slot ${activeTargetSlot}!`);
        }

        function showToast(msg) {
            const toast = document.getElementById('toastMessage');
            document.getElementById('toastText').innerText = msg;
            toast.style.display = 'flex';
            setTimeout(() => { toast.style.display = 'none'; }, 3500);
        }

        window.onload = init;
    </script>
</body>
</html>
"""

class PadiManagerHandler(http.server.SimpleHTTPRequestHandler):
    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path

        if path == "/" or path == "/index.html":
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(HTML_PAGE.encode('utf-8'))
            return

        elif path == "/api/stats":
            conn = get_db()
            cur = conn.cursor()
            cur.execute("SELECT COUNT(*) as total FROM padi_products")
            total = cur.fetchone()['total']
            cur.execute("SELECT batch_num, COUNT(*) as cnt FROM padi_products GROUP BY batch_num")
            batches = {r['batch_num']: r['cnt'] for r in cur.fetchall()}
            conn.close()
            self.send_json({"total": total, "batches": batches})
            return

        elif path == "/api/products":
            conn = get_db()
            cur = conn.cursor()
            cur.execute("""
            SELECT id, batch_num, row_in_batch, kode_sku, nama_produk, harga_satuan, berat_produk_gr,
                   stok_unit, kategori_kode, foto_produk_1, foto_produk_2, foto_produk_3, foto_produk_4, foto_produk_5
            FROM padi_products ORDER BY batch_num ASC, row_in_batch ASC
            """)
            rows = [dict(r) for r in cur.fetchall()]
            conn.close()
            self.send_json(rows)
            return

        elif path.startswith("/api/products/"):
            p_id = path.split("/")[-1]
            conn = get_db()
            cur = conn.cursor()
            cur.execute("SELECT * FROM padi_products WHERE id = ?", (p_id,))
            row = cur.fetchone()
            conn.close()
            if row:
                self.send_json(dict(row))
            else:
                self.send_json({"error": "Product not found"}, status=404)
            return

        elif path == "/api/images/catalog":
            imgs = get_alibaba_images()
            self.send_json(imgs)
            return

        elif path.startswith("/api/download/"):
            b_num = path.split("/")[-1]
            filepath = os.path.join(OUTPUT_PADI, f"PadiUMKM_Bulk_Upload_Batch_{b_num}.xlsx")
            if os.path.exists(filepath):
                with open(filepath, 'rb') as f:
                    content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f"attachment; filename=PadiUMKM_Bulk_Upload_Batch_{b_num}.xlsx")
                self.send_header("Content-Length", str(len(content)))
                self.end_headers()
                self.wfile.write(content)
                return
            else:
                self.send_error(404, "Batch file not found")
                return

        self.send_error(404, "Endpoint not found")

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        length = int(self.headers.get('Content-Length', 0))
        body = self.rfile.read(length).decode('utf-8') if length > 0 else "{}"
        data = json.loads(body) if body else {}

        if path.startswith("/api/products/"):
            p_id = path.split("/")[-1]
            try:
                conn = get_db()
                cur = conn.cursor()
                
                # Fetch current batch and row
                cur.execute("SELECT batch_num, row_in_batch FROM padi_products WHERE id = ?", (p_id,))
                prod = cur.fetchone()
                if not prod:
                    conn.close()
                    self.send_json({"success": False, "error": "Produk tidak ditemukan"}, status=404)
                    return

                batch_num = prod['batch_num']
                row_in_batch = prod['row_in_batch']

                # Update SQLite database
                cur.execute("""
                UPDATE padi_products SET
                    nama_produk = ?, kode_sku = ?, harga_satuan = ?, berat_produk_gr = ?,
                    stok_unit = ?, kategori_kode = ?, dimensi_panjang_cm = ?, dimensi_lebar_cm = ?,
                    dimensi_tinggi_cm = ?, pdn = ?, jenis_tkdn = ?, deskripsi_produk = ?,
                    foto_produk_1 = ?, foto_produk_2 = ?, foto_produk_3 = ?, foto_produk_4 = ?, foto_produk_5 = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """, (
                    data['nama_produk'], data.get('kode_sku', ''), data['harga_satuan'], data['berat_produk_gr'],
                    data['stok_unit'], data['kategori_kode'], data['dimensi_panjang_cm'], data['dimensi_lebar_cm'],
                    data['dimensi_tinggi_cm'], data['pdn'], data['jenis_tkdn'], data['deskripsi_produk'],
                    data.get('foto_produk_1', ''), data.get('foto_produk_2', ''), data.get('foto_produk_3', ''),
                    data.get('foto_produk_4', ''), data.get('foto_produk_5', ''),
                    p_id
                ))
                conn.commit()
                conn.close()

                # Patch Excel batch row immediately!
                patch_excel_row(batch_num, row_in_batch, data)

                self.send_json({"success": True, "message": "Updated in SQLite and Excel"})
            except Exception as e:
                print(f"[ERROR] Failed to update product {p_id}: {e}")
                self.send_json({"success": False, "error": str(e)}, status=500)
            return

        elif path == "/api/export":
            try:
                cnt = export_all_batches_from_db()
                self.send_json({"success": True, "count": cnt})
            except Exception as e:
                print(f"[ERROR] Export failed: {e}")
                self.send_json({"success": False, "error": str(e)}, status=500)
            return

        self.send_error(404, "Endpoint not found")

    def send_json(self, data, status=200):
        body = json.dumps(data).encode('utf-8')
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

def run_server():
    socketserver.TCPServer.allow_reuse_address = True
    active_port = PORT
    httpd = None

    for p in [PORT, 8890, 8891, 8091, 8080]:
        try:
            server_address = ('', p)
            httpd = socketserver.TCPServer(server_address, PadiManagerHandler)
            active_port = p
            break
        except OSError:
            continue

    if not httpd:
        print("[ERROR] Tidak dapat menemukan port yang tersedia untuk PaDi UMKM Manager!")
        return

    print("===================================================================")
    print("       PADI UMKM B2B LISTING MANAGER IS RUNNING LOCALLY!")
    print(f"       URL: http://localhost:{active_port}")
    print(f"       Database: {DB_PATH}")
    print(f"       Excel Batches: {OUTPUT_PADI}")
    print("       Tekan Ctrl + C untuk mematikan server.")
    print("===================================================================")
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n[INFO] PaDi Manager Server stopped.")

if __name__ == '__main__':
    run_server()
